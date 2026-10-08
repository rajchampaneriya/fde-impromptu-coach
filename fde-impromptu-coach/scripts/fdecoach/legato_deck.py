"""Legato practice deck (same blue/Calibri look as the FDE deck) and 1920x1080
Pillow renders of each slide for the YouTube video.

Slides: intro · breath & hum · linking drill · read 1 (phrase map) ·
read 2 (intone) · read 3 (speak it, clean text) · respond (impromptu) · closing.
On the phrase map, underlined words run together as one and "/" marks a
breath ("//" a fuller one at a sentence end).
"""
from __future__ import annotations

import datetime as dt
import logging
from pathlib import Path
from typing import Any, Dict, List, Tuple

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches

from .deck import (BAR_Y, BLUE, BLUE_TINT, CONTENT_W, GRAY, GRAY_LIGHT, MARGIN, NAVY, SLIDE_H, SLIDE_W,
                   _background, _ensure_use_timings, _notes, _rect, _register_notes_master, _spoiler_guard, _text,
                   _timer, fit_font_size, make_show_copy, set_auto_animations, set_transition)
from .legato import marked_tokens, phrase_map
from .library import credit_line

log = logging.getLogger("fdecoach")

HEADER = "LEGATO PRACTICE"
NAMES = ["Breath & hum", "Linking drill", "Read 1 · phrase map", "Read 2 · intone", "Read 3 · speak it", "Respond"]
WARMUP_LINES = [
    "Exhale everything. Breathe in low: belly out, shoulders still.",
    "Hum “mmm” on one breath for a slow count of eight.",
    "Open the hum into “mah”, then glide: may – mee – my – moh – moo.",
    "One breath, one unbroken sound. That is legato.",
]
INTRO_LINES = [
    "Underlined words run together as one word.",
    "Breathe only at  /  — a fuller breath at  //.",
    "Smooth is not fast. Keep the voice moving.",
]
ROUNDS = [  # (label, cue, marked)
    ("Read 1 · phrase map", "One breath per phrase.", True),
    ("Read 2 · intone", "Chant it on one note. No gaps.", True),
    ("Read 3 · speak it", "Natural voice. Keep the thread.", False),
]


def _leg(cfg: Dict[str, Any]) -> Dict[str, Any]:
    return cfg.get("legato", {})


def slide_durations(cfg: Dict[str, Any]) -> List[int]:
    g = _leg(cfg)
    return [int(g.get("intro_seconds", 10)), int(g.get("warmup_seconds", 30)), int(g.get("links_seconds", 45)),
            int(g.get("round1_seconds", 60)), int(g.get("round2_seconds", 60)), int(g.get("round3_seconds", 60)),
            int(g.get("respond_seconds", 45))]


def session_seconds(cfg: Dict[str, Any]) -> int:
    return sum(slide_durations(cfg))


def chapters(session: Dict[str, Any], cfg: Dict[str, Any], offset: float = 0.0) -> List[Tuple[int, str]]:
    durs = slide_durations(cfg)
    marks: List[Tuple[int, str]] = [(0, "Intro")]
    t = durs[0]
    for name, secs in zip(NAMES, durs[1:]):
        marks.append((int(offset + t), name))
        t += secs
    return marks


def credit(session: Dict[str, Any], short: bool = True) -> str:
    return credit_line(session.get("passage") or {}, short=short)


# --------------------------------------------------------------------------- runs

def _marked_runs(text: str, size: int, marked: bool) -> List[Tuple[str, Dict[str, Any]]]:
    if not marked:
        return [(text, {"size": size, "color": NAVY})]
    runs: List[Tuple[str, Dict[str, Any]]] = []
    for t in marked_tokens(text):
        space = " " if t["space"] else ""
        if t["chain"]:
            # the space inside a link is underlined too, so the join reads as one word
            runs.append((t["text"] + (space if t["linked"] else ""),
                         {"size": size, "color": BLUE, "underline": True}))
            if not t["linked"] and not t["mark"]:
                runs.append((space, {"size": size, "color": NAVY}))
        else:
            runs.append((t["text"] + ("" if t["mark"] else space), {"size": size, "color": NAVY}))
        if t["mark"]:
            runs.append((f" {t['mark']} ", {"size": size - 2, "bold": True, "color": GRAY_LIGHT}))
    return [r for r in runs if r[0]]


# --------------------------------------------------------------------------- slides

def _label(s, text: str) -> None:
    _text(s, MARGIN, 0.7, 8.6, 0.35, [(text.upper(), {"size": 14, "bold": True, "color": BLUE, "spacing": 150})])


def _intro_slide(prs, session: Dict[str, Any], stats: Dict[str, Any], cfg: Dict[str, Any]) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _background(s)
    date = dt.date.fromisoformat(session["date"])
    _label(s, f"{HEADER}  ·  CONNECTED SPEECH")
    _text(s, MARGIN, 1.45, 8.0, 1.2, [(f"Day {session['day_number']}", {"size": 66, "color": NAVY})],
          anchor=MSO_ANCHOR.BOTTOM)
    _text(s, MARGIN, 2.75, 8.0, 0.5, [(date.strftime("%A, %d %B %Y"), {"size": 24, "color": GRAY})])
    _text(s, MARGIN, 3.55, 8.2, 1.6, [(line, {"size": 21, "color": NAVY, "newline": True}) for line in INTRO_LINES],
          line_spacing=1.3)
    _text(s, MARGIN, 5.3, 11.5, 0.4, [("Today’s passage:  ", {"size": 15, "bold": True, "color": BLUE}),
                                       (credit(session), {"size": 15, "color": GRAY})])
    px, pw = 9.25, SLIDE_W - MARGIN - 9.25
    _rect(s, px, 1.45, pw, 2.2, BLUE_TINT)
    _text(s, px, 1.7, pw, 0.9, [(str(int(stats.get("current", 0))), {"size": 54, "bold": True, "color": BLUE})],
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    _text(s, px, 2.6, pw, 0.35, [("day streak", {"size": 14, "color": NAVY})], align=PP_ALIGN.CENTER)
    _text(s, px, 3.1, pw, 0.35, [(f"Flow level {session.get('flow_level', 1)} of 3", {"size": 13, "color": GRAY_LIGHT})],
          align=PP_ALIGN.CENTER)
    intro = slide_durations(cfg)[0]
    effects = _timer(s, intro, label=f"Starting in {intro} seconds — sit tall, breathe low", ticks=False)
    set_transition(s, intro * 1000)
    set_auto_animations(s, effects)
    _notes(s, "Audio recording is running. Relax the jaw and shoulders; breathe in through the nose.")


def _warmup_slide(prs, cfg: Dict[str, Any], seconds: int) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _background(s)
    _label(s, "Breath & hum")
    runs = []
    for i, line in enumerate(WARMUP_LINES, 1):
        runs.append((f"{i}   ", {"size": 26, "bold": True, "color": BLUE, "newline": True, "space_before": 14}))
        runs.append((line, {"size": 26, "color": NAVY}))
    _text(s, MARGIN, 1.4, CONTENT_W, 3.9, runs, line_spacing=1.1)
    _text(s, MARGIN, 5.45, CONTENT_W, 0.35, [("Twice through, unhurried.", {"size": 18, "color": GRAY})])
    effects = _timer(s, seconds, label="Breath & hum", ticks=False)
    set_transition(s, seconds * 1000)
    set_auto_animations(s, effects)
    _notes(s, "Legato rides on steady breath. Feel the hum buzz on the lips; keep that buzz when you speak.")


def _links_slide(prs, session: Dict[str, Any], seconds: int) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _background(s)
    _label(s, "Linking drill · from today’s passage")
    items = session.get("chains") or []
    runs: List[Tuple[str, Dict[str, Any]]] = []
    for i, c in enumerate(items[:6]):
        runs.append((c["phrase"], {"size": 30, "color": NAVY, "underline": True, "newline": True,
                                   "space_before": 12 if i else 0}))
        runs.append(("     →  " + c["flow"] if c.get("flow") else "     →  say it as one word",
                     {"size": 24 if c.get("flow") else 18, "bold": bool(c.get("flow")),
                      "color": BLUE if c.get("flow") else GRAY}))
    if not runs:
        runs = [("No joins today: read each phrase on one breath.", {"size": 26, "color": NAVY})]
    _text(s, MARGIN, 1.35, CONTENT_W, 4.0, runs, line_spacing=1.05)
    _text(s, MARGIN, 5.45, CONTENT_W, 0.35, [("Each one twice: slowly with the join, then at speed.",
                                              {"size": 18, "color": GRAY})])
    effects = _timer(s, seconds, label="Linking drill", ticks=False)
    set_transition(s, seconds * 1000)
    set_auto_animations(s, effects)
    _notes(s, "A word ending in a consonant hands that sound to the next word when it starts with a vowel: "
              "turn it off becomes tur-ni-toff. No glottal stop, no gap.")


def _round_slide(prs, session: Dict[str, Any], seconds: int, label: str, cue: str, marked: bool) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _background(s)
    _label(s, label)
    _text(s, SLIDE_W - MARGIN - 4.6, 0.6, 4.6, 0.5, [(cue, {"size": 20, "bold": True, "color": BLUE})],
          align=PP_ALIGN.RIGHT)
    text = session["paragraph"]
    probe = phrase_map(text) if marked else text
    size = fit_font_size(probe, CONTENT_W, 4.4, sizes=(30, 28, 26, 24, 22, 20))
    _text(s, MARGIN, 1.4, CONTENT_W, 4.4, _marked_runs(text, size, marked), line_spacing=1.2, name="Passage")
    _text(s, SLIDE_W - MARGIN - 6.4, BAR_Y - 0.42, 6.4, 0.3, [("— " + credit(session), {"size": 13, "color": GRAY})],
          align=PP_ALIGN.RIGHT, name="Credit")
    effects = _timer(s, seconds, label=label, ticks=False)
    set_transition(s, seconds * 1000)
    set_auto_animations(s, effects)
    _notes(s, f"{cue} Phrase map: {phrase_map(text)}")


def _respond_slide(prs, session: Dict[str, Any], seconds: int) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _background(s)
    _label(s, "Respond · in your own words")
    prompt = session.get("response_prompt", "")
    size = fit_font_size(prompt, CONTENT_W, 2.6, sizes=(40, 38, 36, 34, 32, 30))
    _text(s, MARGIN, 1.5, CONTENT_W, 2.6, [(prompt, {"size": size, "color": NAVY})], anchor=MSO_ANCHOR.MIDDLE,
          line_spacing=1.05, name="Prompt")
    _text(s, MARGIN, 4.35, CONTENT_W, 0.4, [("After: ", {"size": 16, "bold": True, "color": BLUE}),
                                             (credit(session), {"size": 16, "color": GRAY})])
    _text(s, MARGIN, 4.9, CONTENT_W, 0.45, [("Same smooth flow. Pause only where a thought ends.",
                                             {"size": 20, "color": BLUE})])
    cue_at = max(5, seconds - 10)
    effects = _timer(s, seconds, ticks=True, cue=(cue_at, "Wrap up"))
    set_transition(s, seconds * 1000)
    set_auto_animations(s, effects)
    _notes(s, "Impromptu: lead with your point, give one reason or example, land it before the bar runs out — "
              "with the same joined, breath-paced voice you just used on the passage.")


def _closing_slide(prs, session: Dict[str, Any]) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _background(s)
    _label(s, "Session complete")
    _text(s, MARGIN, 1.3, CONTENT_W, 0.9, [("That’s today’s legato.", {"size": 44, "color": NAVY})])
    _text(s, MARGIN, 2.3, CONTENT_W, 0.5, [("Recording stops on its own. Listen back once: find one place your voice "
                                            "stopped inside a phrase.", {"size": 18, "color": GRAY})])
    src = (session.get("passage") or {}).get("source") or {}
    _text(s, MARGIN, 3.3, CONTENT_W, 1.3, [
        ("TODAY’S PASSAGE", {"size": 13, "bold": True, "color": BLUE, "spacing": 150}),
        (credit(session, short=False), {"size": 18, "color": NAVY, "newline": True, "space_before": 6}),
        (f"{src.get('name', '')}  ·  {src.get('url', '')}  ·  public domain",
         {"size": 14, "color": GRAY, "newline": True, "space_before": 4}),
    ], name="Credit")
    set_transition(s, None)


def build_deck(session: Dict[str, Any], stats: Dict[str, Any], cfg: Dict[str, Any], out_dir: Path) -> Tuple[Path, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Emu(int(Inches(SLIDE_W)))
    prs.slide_height = Emu(int(Inches(SLIDE_H)))
    d = slide_durations(cfg)
    _intro_slide(prs, session, stats, cfg)
    _warmup_slide(prs, cfg, d[1])
    _links_slide(prs, session, d[2])
    for (label, cue, marked), secs in zip(ROUNDS, d[3:6]):
        _round_slide(prs, session, secs, label, cue, marked)
    _respond_slide(prs, session, d[6])
    _closing_slide(prs, session)
    _register_notes_master(prs)
    _ensure_use_timings(prs)
    _spoiler_guard(prs)
    cp = prs.core_properties
    cp.title = f"Legato Practice — Day {session['day_number']} ({session['date']})"
    cp.author = "FDE Impromptu Coach"
    stem = f"{session['date']}_Legato"
    pptx_path = out_dir / f"{stem}.pptx"
    tmp = out_dir / f".{stem}.tmp.pptx"
    prs.save(str(tmp))
    tmp.replace(pptx_path)
    ppsx_path = out_dir / f"{stem}.ppsx"
    make_show_copy(pptx_path, ppsx_path)
    return pptx_path, ppsx_path


# --------------------------------------------------------------------------- PNG renders (video)

_W, _H = 1920, 1080
_NAVY, _BLUE, _GRAY, _LIGHT, _SOFT = "#0B2545", "#1F5FBF", "#5B6B7F", "#8A98A9", "#DCE7F7"


def _draw_marked(draw, get, text: str, marked: bool, xy: Tuple[int, int], wh: Tuple[int, int]) -> None:
    """Wrapped passage; underlines joins (including the joined space); grey
    breath marks. Original spacing is kept ("men—that" stays joined)."""
    pieces: List[Tuple[str, str, bool, bool]] = []  # (text, kind, joined_to_next, space_after)
    for t in marked_tokens(text):
        kind = "chain" if (marked and t["chain"]) else "plain"
        has_mark = marked and bool(t["mark"])
        pieces.append((t["text"], kind, marked and t["linked"], t["space"] or has_mark))
        if has_mark:
            pieces.append((t["mark"], "mark", False, t["space"] or True))
    size = 46
    while size >= 26:
        font = get(size)
        space = draw.textlength(" ", font=font)
        lines: List[List[Tuple[str, str, bool, bool]]] = [[]]
        width = 0.0
        for p in pieces:
            w = draw.textlength(p[0], font=font) + (space if p[3] else 0)
            if lines[-1] and width + w > wh[0] and lines[-1][-1][3]:
                lines.append([])
                width = 0.0
            lines[-1].append(p)
            width += w
        if len(lines) * size * 1.45 <= wh[1]:
            break
        size -= 2
    font = get(size)
    space = draw.textlength(" ", font=font)
    x0, y = xy
    for line in lines:
        x = x0
        for i, (word, kind, joined, space_after) in enumerate(line):
            color = {"chain": _BLUE, "mark": _LIGHT, "plain": _NAVY}[kind]
            draw.text((x, y), word, font=font, fill=color)
            w = draw.textlength(word, font=font)
            gap = space if space_after else 0
            if kind == "chain":
                end = x + w + (gap if joined and i < len(line) - 1 else 0)
                draw.line((x, y + size * 1.12, end, y + size * 1.12), fill=_BLUE, width=max(2, size // 18))
            x += w + gap
        y += int(size * 1.45)


def render_slide_pngs(session: Dict[str, Any], cfg: Dict[str, Any], out_dir: Path) -> List[Path]:
    """One 1920x1080 PNG per slide (intro..closing), in order."""
    try:
        from PIL import Image, ImageDraw
    except ImportError:
        log.warning("Pillow not installed; cannot render legato slides")
        return []
    from .pronounce_deck import _fonts
    out_dir.mkdir(parents=True, exist_ok=True)
    get = _fonts()
    date = dt.date.fromisoformat(session["date"])
    paths: List[Path] = []

    def slide():
        img = Image.new("RGB", (_W, _H), "white")
        return img, ImageDraw.Draw(img)

    def save(img, idx: int, name: str) -> None:
        p = out_dir / f"slide{idx}_{name}.png"
        img.save(p, "PNG")
        paths.append(p)

    img, dr = slide()
    dr.text((128, 90), f"{HEADER}  ·  CONNECTED SPEECH", font=get(36, True), fill=_BLUE)
    dr.text((128, 170), f"Day {session['day_number']}", font=get(150, True), fill=_NAVY)
    dr.text((128, 400), date.strftime("%A, %d %B %Y"), font=get(56), fill=_GRAY)
    for i, line in enumerate(INTRO_LINES):
        dr.text((128, 520 + i * 70), line, font=get(44, True), fill=_NAVY)
    dr.text((128, 800), "Today’s passage: " + credit(session), font=get(34), fill=_GRAY)
    save(img, 1, "intro")

    img, dr = slide()
    dr.text((128, 80), "BREATH & HUM", font=get(34, True), fill=_BLUE)
    for i, line in enumerate(WARMUP_LINES):
        dr.text((128, 200 + i * 130), f"{i + 1}", font=get(56, True), fill=_BLUE)
        dr.text((210, 200 + i * 130), line, font=get(52), fill=_NAVY)
    save(img, 2, "warmup")

    img, dr = slide()
    dr.text((128, 80), "LINKING DRILL · FROM TODAY’S PASSAGE", font=get(34, True), fill=_BLUE)
    y = 190
    for c in (session.get("chains") or [])[:6]:
        f = get(60, True)
        dr.text((128, y), c["phrase"], font=f, fill=_NAVY)
        w = dr.textlength(c["phrase"], font=f)
        dr.line((128, y + 72, 128 + w, y + 72), fill=_BLUE, width=4)
        dr.text((128 + w + 40, y + 6), "→  " + (c.get("flow") or "say it as one word"),
                font=get(48 if c.get("flow") else 36, bool(c.get("flow"))), fill=_BLUE if c.get("flow") else _GRAY)
        y += 128
    dr.text((128, 980), "Each one twice: slowly with the join, then at speed.", font=get(38), fill=_GRAY)
    save(img, 3, "links")

    for idx, (label, cue, marked) in enumerate(ROUNDS, start=4):
        img, dr = slide()
        dr.text((128, 80), label.upper(), font=get(34, True), fill=_BLUE)
        w = dr.textlength(cue, font=get(44, True))
        dr.text((_W - 128 - w, 72), cue, font=get(44, True), fill=_BLUE)
        _draw_marked(dr, get, session["paragraph"], marked, (128, 200), (1664, 730))
        line = "— " + credit(session)
        dr.text((_W - 128 - dr.textlength(line, font=get(30)), 950), line, font=get(30), fill=_GRAY)
        dr.rectangle((128, 1000, _W - 128, 1020), fill=_SOFT)
        save(img, idx, f"read{idx - 3}")

    img, dr = slide()
    dr.text((128, 80), "RESPOND · IN YOUR OWN WORDS", font=get(34, True), fill=_BLUE)
    words, line, y = session.get("response_prompt", "").split(), "", 260
    f = get(80)
    for word in words:  # simple wrap for the big prompt
        trial = (line + " " + word).strip()
        if dr.textlength(trial, font=f) > 1664:
            dr.text((128, y), line, font=f, fill=_NAVY)
            y, line = y + 100, word
        else:
            line = trial
    dr.text((128, y), line, font=f, fill=_NAVY)
    dr.text((128, y + 150), "After: " + credit(session), font=get(38), fill=_GRAY)
    dr.text((128, y + 220), "Same smooth flow. Pause only where a thought ends.", font=get(42, True), fill=_BLUE)
    save(img, 7, "respond")

    img, dr = slide()
    src = (session.get("passage") or {}).get("source") or {}
    dr.text((128, 90), "SESSION COMPLETE", font=get(34, True), fill=_BLUE)
    dr.text((128, 200), "That’s today’s legato.", font=get(90, True), fill=_NAVY)
    dr.text((128, 360), "Listen back once: find one place your voice stopped inside a phrase.", font=get(42),
            fill=_GRAY)
    dr.text((128, 520), "TODAY’S PASSAGE", font=get(32, True), fill=_BLUE)
    dr.text((128, 580), credit(session, short=False), font=get(40), fill=_NAVY)
    dr.text((128, 650), f"{src.get('name', '')}  ·  {src.get('url', '')}  ·  public domain", font=get(32), fill=_GRAY)
    save(img, 8, "closing")
    return paths
