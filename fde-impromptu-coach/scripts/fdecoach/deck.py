"""Build the daily practice deck with python-pptx.

Design: 16:9, white background, navy text, one blue accent, Calibri only.
Each question slide has a 60-second progress bar (linear wipe animation that
starts automatically) and auto-advances after 60 seconds. PowerPoint waits for
automatic animations to finish before auto-advancing, so bar and slide change
stay in sync. Only plain shapes, text boxes and notes are used so the file
imports cleanly into Google Slides.
"""
from __future__ import annotations

import datetime as dt
import math
import shutil
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Sequence, Tuple

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Inches, Pt

FONT = "Calibri"
NAVY = RGBColor(0x0B, 0x25, 0x45)
BLUE = RGBColor(0x1F, 0x5F, 0xBF)
BLUE_SOFT = RGBColor(0xDC, 0xE7, 0xF7)
BLUE_TINT = RGBColor(0xF1, 0xF6, 0xFD)
GRAY = RGBColor(0x5B, 0x6B, 0x7F)
GRAY_LIGHT = RGBColor(0x8A, 0x98, 0xA9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.85
CONTENT_W = SLIDE_W - 2 * MARGIN
BAR_Y, BAR_H = 6.42, 0.11

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PPSX_CT = "application/vnd.openxmlformats-officedocument.presentationml.slideshow.main+xml"
PPTX_CT = "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"


# --------------------------------------------------------------------------- primitives

def _text(slide, x: float, y: float, w: float, h: float, runs: Sequence[Tuple[str, Dict[str, Any]]],
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing: float = 1.0, name: str = ""):
    """runs: [(text, {size, color, bold, italic, spacing, newline})]; newline starts a new paragraph."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    para = tf.paragraphs[0]
    para.alignment = align
    para.line_spacing = line_spacing
    first = True
    for text, style in runs:
        if style.get("newline") and not first:
            para = tf.add_paragraph()
            para.alignment = align
            para.line_spacing = line_spacing
            if style.get("space_before"):
                para.space_before = Pt(style["space_before"])
        first = False
        run = para.add_run()
        run.text = text
        f = run.font
        f.name = FONT
        f.size = Pt(style.get("size", 18))
        f.bold = bool(style.get("bold", False))
        f.italic = bool(style.get("italic", False))
        f.color.rgb = style.get("color", NAVY)
        if style.get("spacing"):
            run._r.get_or_add_rPr().set("spc", str(int(style["spacing"])))
    return box


def _rect(slide, x: float, y: float, w: float, h: float, color: RGBColor, shape=MSO_SHAPE.RECTANGLE, name: str = ""):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        shp.name = name
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    # Drop the theme style reference so no renderer (Google Slides, Keynote, LibreOffice) adds a shadow.
    style = shp._element.find(f"{{{P_NS}}}style")
    if style is not None:
        shp._element.remove(style)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = 0.08
    return shp


def _background(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE


def fit_font_size(text: str, width_in: float, height_in: float,
                  sizes: Sequence[int] = (40, 38, 36, 34, 32, 30, 28, 26, 24, 22)) -> int:
    """Largest size whose simulated word-wrap fits the box (Calibri avg glyph ~0.5 em)."""
    words = text.split()
    for size in sizes:
        chars_per_line = max(8, int(width_in * 72 / (size * 0.50)))
        lines, current = 1, 0
        for word in words:
            need = len(word) + (1 if current else 0)
            if current + need > chars_per_line:
                lines += 1
                current = len(word)
            else:
                current += need
        if lines * size * 1.2 <= height_in * 72 * 0.92:
            return size
    return sizes[-1]


# --------------------------------------------------------------------------- slide timing XML

def set_transition(slide, advance_ms: "int | None") -> None:
    adv = f' advTm="{int(advance_ms)}"' if advance_ms else ""
    el = etree.fromstring(f'<p:transition xmlns:p="{P_NS}" spd="med"{adv}><p:fade/></p:transition>')
    _insert_after_clrmap(slide, el)


def _insert_after_clrmap(slide, el) -> None:
    sld = slide._element
    for tag in ("transition", "timing"):
        for old in sld.findall(f"{{{P_NS}}}{tag}"):
            if old.tag == el.tag:
                sld.remove(old)
    clr = sld.find(f"{{{P_NS}}}clrMapOvr")
    anchor = clr if clr is not None else sld.find(f"{{{P_NS}}}cSld")
    if el.tag == f"{{{P_NS}}}timing":
        tr = sld.find(f"{{{P_NS}}}transition")
        anchor = tr if tr is not None else anchor
    anchor.addnext(el)


def set_auto_animations(slide, effects: List[Dict[str, Any]]) -> None:
    """effects: [{'spid': int, 'kind': 'wipe'|'appear', 'delay': ms, 'dur': ms}]
    All effects start automatically when the slide appears ("With Previous" + delay)."""
    ids = iter(range(5, 1000))
    parts: List[str] = []
    for eff in effects:
        spid = int(eff["spid"])
        ctn = next(ids)
        set_id = next(ids)
        set_xml = (
            f'<p:set><p:cBhvr><p:cTn id="{set_id}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/>'
            f'</p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst>'
            f'<p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/>'
            f'</p:to></p:set>'
        )
        if eff["kind"] == "wipe":
            anim_id = next(ids)
            body = set_xml + (
                f'<p:animEffect transition="in" filter="wipe(left)"><p:cBhvr><p:cTn id="{anim_id}" '
                f'dur="{int(eff["dur"])}"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
            )
            preset = 'presetID="22" presetClass="entr" presetSubtype="8"'
        else:
            body = set_xml
            preset = 'presetID="1" presetClass="entr" presetSubtype="0"'
        parts.append(
            f'<p:par><p:cTn id="{ctn}" {preset} fill="hold" grpId="0" nodeType="withEffect">'
            f'<p:stCondLst><p:cond delay="{int(eff.get("delay", 0))}"/></p:stCondLst>'
            f'<p:childTnLst>{body}</p:childTnLst></p:cTn></p:par>'
        )
    bld = "".join(f'<p:bldP spid="{int(e["spid"])}" grpId="0" animBg="1"/>' for e in effects)
    xml = (
        f'<p:timing xmlns:p="{P_NS}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">'
        '<p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq">'
        '<p:childTnLst><p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/>'
        '<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
        '<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
        + "".join(parts) +
        '</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn>'
        '<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
        '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
        '</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>'
        f'<p:bldLst>{bld}</p:bldLst></p:timing>'
    )
    _insert_after_clrmap(slide, etree.fromstring(xml))


# --------------------------------------------------------------------------- components

def _timer(slide, seconds: int, label: str = "", ticks: bool = True, cue: "Tuple[int, str] | None" = None) -> List[Dict[str, Any]]:
    """Track + animated bar (+ optional ticks and a delayed cue). Returns animation effects."""
    _rect(slide, MARGIN, BAR_Y, CONTENT_W, BAR_H, BLUE_SOFT, name="Timer track")
    bar = _rect(slide, MARGIN, BAR_Y, CONTENT_W, BAR_H, BLUE, name="Timer bar")
    effects = [{"spid": bar.shape_id, "kind": "wipe", "delay": 0, "dur": seconds * 1000}]
    if ticks:
        marks = [s for s in (15, 30, 45) if s < seconds] + [seconds]
        for s in marks:
            cx = MARGIN + CONTENT_W * s / seconds
            _rect(slide, cx - 0.01, BAR_Y - 0.05, 0.02, BAR_H + 0.1, GRAY_LIGHT if s != seconds else BLUE_SOFT,
                  name=f"Tick {s}s")
            mm, ss = divmod(s, 60)
            last = s == seconds
            _text(slide, cx - (1.0 if last else 0.5), BAR_Y + 0.2, 1.0, 0.3,
                  [(f"{mm}:{ss:02d}", {"size": 11, "color": GRAY_LIGHT})],
                  align=PP_ALIGN.RIGHT if last else PP_ALIGN.CENTER)
    if label:
        _text(slide, MARGIN, BAR_Y - 0.42, 6.0, 0.3, [(label, {"size": 13, "color": GRAY})])
    if cue:
        at, text = cue
        box = _text(slide, SLIDE_W - MARGIN - 3.0, BAR_Y - 0.5, 3.0, 0.4,
                    [(text, {"size": 18, "bold": True, "color": BLUE})], align=PP_ALIGN.RIGHT,
                    anchor=MSO_ANCHOR.BOTTOM, name="Wrap-up cue")
        effects.append({"spid": box.shape_id, "kind": "appear", "delay": at * 1000, "dur": 0})
    return effects


def _difficulty_dots(slide, level: int) -> None:
    d, gap = 0.16, 0.26
    right = SLIDE_W - MARGIN
    x0 = right - (4 * gap + d)
    y = 0.74
    _text(slide, x0 - 1.75, 0.68, 1.6, 0.3, [("DIFFICULTY", {"size": 11, "color": GRAY_LIGHT, "spacing": 150})],
          align=PP_ALIGN.RIGHT)
    for i in range(5):
        _rect(slide, x0 + i * gap, y, d, d, BLUE if i < level else BLUE_SOFT, shape=MSO_SHAPE.OVAL,
              name=f"Difficulty dot {i + 1}")


def _notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


# --------------------------------------------------------------------------- slides

def _intro_slide(prs, session: Dict[str, Any], stats: Dict[str, Any], cfg: Dict[str, Any]) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _background(s)
    date = dt.date.fromisoformat(session["date"])
    n = len(session["questions"])
    secs = int(cfg.get("seconds_per_question", 60))
    _text(s, MARGIN, 0.7, 9.0, 0.35, [(f"{cfg.get('role', 'Forward Deployed Engineer').upper()}  \u00b7  IMPROMPTU PRACTICE",
                                        {"size": 14, "bold": True, "color": BLUE, "spacing": 150})])
    _text(s, MARGIN, 1.45, 8.0, 1.2, [(f"Day {session['day_number']}", {"size": 66, "color": NAVY})],
          anchor=MSO_ANCHOR.BOTTOM)
    _text(s, MARGIN, 2.75, 8.0, 0.5, [(date.strftime("%A, %d %B %Y"), {"size": 24, "color": GRAY})])
    _text(s, MARGIN, 3.7, 8.0, 0.45, [(f"{n} questions  \u00b7  {secs} seconds each  \u00b7  no preparation",
                                        {"size": 22, "color": NAVY})])
    _text(s, MARGIN, 4.3, 7.8, 0.8, [("Pause. Lead with your point. Land it before the bar runs out.",
                                       {"size": 18, "color": GRAY})])
    # streak panel
    px, py, pw, ph = 9.25, 1.45, SLIDE_W - MARGIN - 9.25, 2.9
    _rect(s, px, py, pw, ph, BLUE_TINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name="Streak panel")
    current = int(stats.get("current", 0))
    _text(s, px, py + 0.3, pw, 1.1, [(str(current), {"size": 60, "bold": True, "color": BLUE})],
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    _text(s, px, py + 1.45, pw, 0.4, [("day streak", {"size": 16, "color": NAVY})], align=PP_ALIGN.CENTER)
    sub = f"Today makes it {current + 1}" if current else (
        "New streak starts today" if stats.get("total") else "Your first session")
    _text(s, px, py + 1.9, pw, 0.35, [(sub, {"size": 13, "color": GRAY})], align=PP_ALIGN.CENTER)
    _text(s, px, py + 2.3, pw, 0.35, [(f"Best {stats.get('best', 0)}  \u00b7  Level {session['level']} of 5",
                                        {"size": 13, "color": GRAY_LIGHT})], align=PP_ALIGN.CENTER)
    intro = int(cfg.get("intro_seconds", 10))
    effects = _timer(s, intro, label=f"Starting in {intro} seconds \u2014 sit tall, look at the camera", ticks=False)
    set_transition(s, intro * 1000)
    set_auto_animations(s, effects)
    _notes(s, "Recording is running. Use these seconds to settle: shoulders down, one slow breath, eyes on the camera.")


def _question_slide(prs, q: Dict[str, Any], idx: int, total: int, level: int, cfg: Dict[str, Any]) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _background(s)
    secs = int(cfg.get("seconds_per_question", 60))
    _text(s, MARGIN, 0.7, 8.6, 0.35, [
        (f"QUESTION {idx} OF {total}", {"size": 14, "bold": True, "color": BLUE, "spacing": 150}),
        (f"   \u00b7   {q.get('category_label', '').upper()}", {"size": 14, "color": GRAY, "spacing": 100}),
    ])
    _difficulty_dots(s, int(q.get("difficulty", 1)))
    has_constraint = bool(q.get("constraint"))
    q_top, q_h = 1.45, (3.45 if has_constraint else 4.0)
    size = fit_font_size(q["text"], CONTENT_W, q_h)
    _text(s, MARGIN, q_top, CONTENT_W, q_h, [(q["text"], {"size": size, "color": NAVY})],
          anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05, name="Question")
    if has_constraint:
        _text(s, MARGIN, 5.05, CONTENT_W, 0.45, [
            ("CONSTRAINT  ", {"size": 13, "bold": True, "color": BLUE, "spacing": 150}),
            (q["constraint"], {"size": 20, "color": BLUE}),
        ], anchor=MSO_ANCHOR.MIDDLE, name="Constraint")
    if level <= int(cfg.get("show_framework_hint_until_level", 2)):
        _text(s, MARGIN, 5.62, 7.6, 0.32, [("Try:  " + q.get("framework", ""), {"size": 13, "color": GRAY_LIGHT})],
              name="Framework hint")
    cue_at = int(cfg.get("wrap_up_cue_seconds", 45))
    effects = _timer(s, secs, ticks=True, cue=(cue_at, "Wrap up") if 0 < cue_at < secs else None)
    set_transition(s, secs * 1000)
    set_auto_animations(s, effects)
    notes = [
        f"Competency: {q.get('category_label', '')}   |   Difficulty {q.get('difficulty')}/5   |   Format: {q.get('format')}",
    ]
    if q.get("constraint"):
        notes.append(f"Constraint: {q['constraint']}")
    notes += [
        "",
        f"What a strong answer does: {q.get('coach_note', '')}",
        f"Framework: {q.get('framework', '')}",
        "",
        "Self-review when you watch the video:",
        "[ ] Clear point in the first 10 seconds",
        "[ ] One concrete example, number or story",
        "[ ] Calm pace, deliberate pauses, few fillers",
        "[ ] Ended cleanly (no trailing \"so yeah\") before the bar ran out",
        "[ ] Would this listener trust me more after hearing it?",
    ]
    _notes(s, "\n".join(notes))


def _closing_slide(prs, session: Dict[str, Any]) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _background(s)
    _text(s, MARGIN, 0.7, 8.0, 0.35, [("SESSION COMPLETE", {"size": 14, "bold": True, "color": BLUE, "spacing": 150})])
    words = {3: "three", 4: "four", 5: "five", 6: "six", 7: "seven"}
    count = words.get(len(session["questions"]), str(len(session["questions"])))
    _text(s, MARGIN, 1.2, CONTENT_W, 0.9, [(f"That\u2019s today\u2019s {count}.", {"size": 44, "color": NAVY})])
    _text(s, MARGIN, 2.15, CONTENT_W, 0.5, [(
        "Recording stops on its own. Tonight, watch once: note one thing to keep and one thing to fix.",
        {"size": 18, "color": GRAY})])
    runs: List[Tuple[str, Dict[str, Any]]] = []
    for i, q in enumerate(session["questions"], 1):
        text = q["text"] if len(q["text"]) <= 175 else q["text"][:172].rstrip() + "\u2026"
        runs.append((f"{i}.  ", {"size": 15, "bold": True, "color": BLUE, "newline": True, "space_before": 8 if i > 1 else 0}))
        runs.append((text, {"size": 15, "color": NAVY}))
    _text(s, MARGIN, 3.05, CONTENT_W, 3.9, runs, line_spacing=1.05, name="Today's questions")
    set_transition(s, None)


# --------------------------------------------------------------------------- build

def _register_notes_master(prs) -> None:
    """python-pptx's template omits <p:notesMasterIdLst> from presentation.xml even
    though the notes master part and its relationship exist. PowerPoint tolerates the
    omission but Keynote refuses to import the whole file, so register it explicitly."""
    if not any(getattr(s, "has_notes_slide", False) for s in prs.slides):
        return
    root = prs.part._element
    if root.find(f"{{{P_NS}}}notesMasterIdLst") is not None:
        return
    rid = next((rel.rId for rel in prs.part.rels.values() if rel.reltype.endswith("/notesMaster")), None)
    if rid is None:
        return
    nml = etree.Element(f"{{{P_NS}}}notesMasterIdLst")
    mid = etree.SubElement(nml, f"{{{P_NS}}}notesMasterId")
    mid.set(f"{{{R_NS}}}id", rid)
    anchor = root.find(f"{{{P_NS}}}sldMasterIdLst")
    if anchor is not None:
        anchor.addnext(nml)
    else:
        root.insert(0, nml)


def _ensure_use_timings(prs) -> None:
    """Explicitly enable 'Use Timings' so auto-advance works even if a user default differs."""
    part = None
    for rel in prs.part.rels.values():
        if rel.reltype.endswith("/presProps"):
            part = rel.target_part
    if part is None:
        return
    root = etree.fromstring(part.blob)
    show = root.find(f"{{{P_NS}}}showPr")
    if show is None:
        show = etree.Element(f"{{{P_NS}}}showPr")
        successors = [root.find(f"{{{P_NS}}}{t}") for t in ("clrMru", "extLst")]
        successors = [e for e in successors if e is not None]
        if successors:
            successors[0].addprevious(show)
        else:
            root.append(show)
    show.set("useTimings", "1")
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _spoiler_guard(prs) -> None:
    """Open in Normal view with the slide-thumbnail pane collapsed, so opening the deck
    before recording shows only the intro slide instead of spoiling the questions."""
    for rel in prs.part.rels.values():
        if not rel.reltype.endswith("/viewProps"):
            continue
        part = rel.target_part
        root = etree.fromstring(part.blob)
        root.attrib.pop("lastView", None)  # default = sldView (Normal view), not the thumbnail sorter
        nv = root.find(f"{{{P_NS}}}normalViewPr")
        if nv is None:
            nv = etree.Element(f"{{{P_NS}}}normalViewPr")
            root.insert(0, nv)
        nv.set("snapVertSplitter", "1")
        nv.set("vertBarState", "minimized")
        left = nv.find(f"{{{P_NS}}}restoredLeft")
        if left is not None:
            left.set("sz", "5000")
            left.set("autoAdjust", "0")
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def build_deck(session: Dict[str, Any], stats: Dict[str, Any], cfg: Dict[str, Any], out_dir: Path) -> Tuple[Path, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Emu(int(Inches(SLIDE_W)))
    prs.slide_height = Emu(int(Inches(SLIDE_H)))
    total = len(session["questions"])
    _intro_slide(prs, session, stats, cfg)
    for i, q in enumerate(session["questions"], 1):
        _question_slide(prs, q, i, total, int(session["level"]), cfg)
    _closing_slide(prs, session)
    _register_notes_master(prs)
    _ensure_use_timings(prs)
    _spoiler_guard(prs)
    cp = prs.core_properties
    cp.title = f"FDE Impromptu Practice \u2014 Day {session['day_number']} ({session['date']})"
    cp.author = "FDE Impromptu Coach"
    cp.subject = "Daily impromptu communication practice"
    cp.keywords = "impromptu; communication; leadership; forward deployed engineer"
    stem = f"{session['date']}_FDE-Impromptu"
    pptx_path = out_dir / f"{stem}.pptx"
    tmp = out_dir / f".{stem}.tmp.pptx"
    prs.save(str(tmp))
    tmp.replace(pptx_path)
    ppsx_path = out_dir / f"{stem}.ppsx"
    make_show_copy(pptx_path, ppsx_path)
    return pptx_path, ppsx_path


def make_show_copy(pptx_path: Path, ppsx_path: Path) -> None:
    """A .ppsx opens straight into slide show in PowerPoint (used as a fallback start method)."""
    tmp = ppsx_path.with_name("." + ppsx_path.name + ".tmp")
    with zipfile.ZipFile(pptx_path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(PPTX_CT.encode(), PPSX_CT.encode())
            dst.writestr(item, data)
    shutil.move(str(tmp), str(ppsx_path))


def session_seconds(cfg: Dict[str, Any], n_questions: int) -> int:
    return int(cfg.get("intro_seconds", 10)) + n_questions * int(cfg.get("seconds_per_question", 60))


def chapters(session: Dict[str, Any], cfg: Dict[str, Any], offset: float = 0.0) -> List[Tuple[int, str]]:
    """YouTube chapter marks (seconds, title) for an auto-synced recording."""
    intro = int(cfg.get("intro_seconds", 10))
    per = int(cfg.get("seconds_per_question", 60))
    marks = [(0, "Intro")]
    for i, q in enumerate(session["questions"], 1):
        start = int(math.floor(offset + intro + (i - 1) * per))
        marks.append((start, f"Q{i} \u00b7 {q.get('category_label', '')}"))
    marks.append((int(math.floor(offset + intro + len(session['questions']) * per)), "Wrap-up"))
    return marks
