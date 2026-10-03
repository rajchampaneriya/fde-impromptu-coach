"""Software Factory slides: one Pillow renderer for every frame (1920x1080), used
both by the take deck you record with and by the final video, so what you see
while speaking is exactly what viewers see.

Design: white, navy text, one blue accent, Calibri (from Microsoft Office when
installed), one idea per slide, a picture over words. The take deck is a .pptx
of full-slide images plus the same animated progress bar and auto-advance
timings as the impromptu deck.
"""
from __future__ import annotations

import logging
import math
import os
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

from .factory import DIAGRAMS_DIR, TAKE_FOCUS, module_info, segments

log = logging.getLogger("fdecoach")

W, H = 1920, 1080
M = 120
NAVY, BLUE, SOFT, TINT = "#0B2545", "#1F5FBF", "#DCE7F7", "#F1F6FD"
GRAY, GRAY_L, NEUTRAL, WHITE = "#5B6B7F", "#8A98A9", "#F4F6F9", "#FFFFFF"
CODE_FG, CODE_DIM, CODE_PROMPT = "#E6EEF9", "#9FB3CF", "#8FBCFF"
CONTENT_BOTTOM = 900  # the take deck's progress bar sits at ~924 px
TRACKER = (("hook", "WHAT & WHY"), ("demo", "SEE IT"), ("picture", "GROUP IT"), ("end", "TAKEAWAY"))
TERM_BG, TERM_BAR, TERM_OUT, TERM_HI, AMBER = "#0B2545", "#13325C", "#C9D6EA", "#1E4E8C", "#F2B84B"

_OFFICE_FONT_DIRS = [
    "/Applications/Microsoft PowerPoint.app/Contents/Resources/DFonts",
    "/Applications/Microsoft Word.app/Contents/Resources/DFonts",
    "/Applications/Microsoft Excel.app/Contents/Resources/DFonts",
    "/Applications/Microsoft Outlook.app/Contents/Resources/DFonts",
    "~/Library/Fonts", "/Library/Fonts",
]
_NAMES = {
    "regular": ["calibri.ttf", "calibri regular.ttf"],
    "bold": ["calibrib.ttf", "calibri bold.ttf", "calibri-bold.ttf"],
}
_FALLBACKS = {
    "regular": ["/usr/share/fonts/truetype/crosextra/Carlito-Regular.ttf",
                "/System/Library/Fonts/Supplemental/Arial.ttf", "/Library/Fonts/Arial.ttf",
                "/System/Library/Fonts/Helvetica.ttc", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"],
    "bold": ["/usr/share/fonts/truetype/crosextra/Carlito-Bold.ttf",
             "/System/Library/Fonts/Supplemental/Arial Bold.ttf", "/Library/Fonts/Arial Bold.ttf",
             "/System/Library/Fonts/Helvetica.ttc", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"],
    "mono": ["/System/Library/Fonts/Menlo.ttc", "/System/Library/Fonts/SFNSMono.ttf",
             "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf"],
}
_FONT_PATHS: Dict[str, Optional[str]] = {}
_FONT_CACHE: Dict[Tuple[str, int], Any] = {}


# --------------------------------------------------------------------------- fonts

def _find_office(style: str) -> Optional[str]:
    for base in _OFFICE_FONT_DIRS:
        folder = Path(base).expanduser()
        try:
            entries = {p.name.lower(): p for p in folder.iterdir()}
        except OSError:
            continue
        for name in _NAMES[style]:
            if name in entries:
                return str(entries[name])
    return None


def resolve_fonts(cfg: Optional[Dict[str, Any]] = None) -> Dict[str, Optional[str]]:
    """Font files for regular / bold / mono: configured path, then Calibri from Office, then
    metric-compatible Carlito, then Arial/Helvetica (macOS) or DejaVu (Linux)."""
    if _FONT_PATHS:
        return _FONT_PATHS
    f = (cfg or {}).get("factory", {})
    for style in ("regular", "bold", "mono"):
        chosen = None
        configured = f.get(f"font_{style}") if style != "mono" else None
        if configured and Path(configured).expanduser().exists():
            chosen = str(Path(configured).expanduser())
        if not chosen and style in _NAMES and os.environ.get("FDE_COACH_NO_SYSTEM_FONTS") != "1":
            chosen = _find_office(style)
        if not chosen:
            chosen = next((p for p in _FALLBACKS[style] if Path(p).exists()), None)
        _FONT_PATHS[style] = chosen
    return _FONT_PATHS


def font_report(cfg: Dict[str, Any]) -> str:
    fonts = resolve_fonts(cfg)
    reg = fonts.get("regular") or "Pillow default"
    calibri = "calibri" in reg.lower()
    return f"{'Calibri' if calibri else Path(reg).stem + ' (Calibri not found)'} — {reg}"


def font(size: int, style: str = "regular"):
    from PIL import ImageFont
    key = (style, int(size))
    if key not in _FONT_CACHE:
        path = resolve_fonts().get(style)
        try:
            _FONT_CACHE[key] = ImageFont.truetype(path, int(size)) if path else ImageFont.load_default()
        except OSError:
            _FONT_CACHE[key] = ImageFont.load_default()
    return _FONT_CACHE[key]


# --------------------------------------------------------------------------- drawing helpers

def _tw(draw, text: str, f) -> float:
    return draw.textlength(text, font=f)


def wrap(draw, text: str, f, max_w: float) -> List[str]:
    lines: List[str] = []
    for para in str(text).split("\n"):
        line = ""
        for word in para.split():
            cand = f"{line} {word}".strip()
            if line and _tw(draw, cand, f) > max_w:
                lines.append(line)
                line = word
            else:
                line = cand
        lines.append(line)
    return lines


def fit(draw, text: str, max_w: float, max_h: float, sizes: Sequence[int], style: str = "regular",
        spacing: float = 1.22, max_lines: int = 0):
    """Largest size whose wrapped text fits the box. Returns (font, lines, line_height)."""
    for size in sizes:
        f = font(size, style)
        lines = wrap(draw, text, f, max_w)
        lh = int(size * spacing)
        if len(lines) * lh <= max_h and (not max_lines or len(lines) <= max_lines) \
                and all(_tw(draw, ln, f) <= max_w for ln in lines):
            return f, lines, lh
    f = font(sizes[-1], style)
    return f, wrap(draw, text, f, max_w), int(sizes[-1] * spacing)


def draw_lines(draw, x: float, y: float, lines: Sequence[str], f, lh: int, fill: str,
               align: str = "left", width: float = 0) -> float:
    for ln in lines:
        lx = x
        if align == "center":
            lx = x + (width - _tw(draw, ln, f)) / 2
        elif align == "right":
            lx = x + width - _tw(draw, ln, f)
        draw.text((lx, y), ln, font=f, fill=fill)
        y += lh
    return y


def spaced(text: str) -> str:
    """Letter-spaced caps for small labels (Pillow has no tracking)."""
    return " ".join(text.upper())


def rounded(draw, box, fill: str, outline: Optional[str] = None, width: int = 0, radius: int = 26) -> None:
    draw.rounded_rectangle(tuple(int(v) for v in box), radius=radius, fill=fill, outline=outline, width=width)


def arrow(draw, p0, p1, color: str = BLUE, width: int = 6, head: int = 22) -> None:
    (x0, y0), (x1, y1) = p0, p1
    ang = math.atan2(y1 - y0, x1 - x0)
    bx, by = x1 - head * math.cos(ang), y1 - head * math.sin(ang)
    draw.line((x0, y0, bx, by), fill=color, width=width)
    left = (bx + head * 0.6 * math.cos(ang + math.pi / 2), by + head * 0.6 * math.sin(ang + math.pi / 2))
    right = (bx + head * 0.6 * math.cos(ang - math.pi / 2), by + head * 0.6 * math.sin(ang - math.pi / 2))
    draw.polygon([(x1, y1), left, right], fill=color)


def _canvas():
    from PIL import Image, ImageDraw
    img = Image.new("RGB", (W, H), WHITE)
    return img, ImageDraw.Draw(img)


# --------------------------------------------------------------------------- slide chrome

def _label(draw, text: str) -> None:
    draw.text((M, 64), spaced(text), font=font(34, "bold"), fill=BLUE)


def _tracker(draw, current: str) -> None:
    sep = "   ·   "
    f_on, f_off = font(28, "bold"), font(28)
    parts = [(lbl, f_on if key == current else f_off, BLUE if key == current else GRAY_L)
             for key, lbl in TRACKER]
    total = sum(_tw(draw, p[0], p[1]) for p in parts) + _tw(draw, sep, f_off) * (len(parts) - 1)
    x = W - M - total
    for i, (lbl, f, col) in enumerate(parts):
        draw.text((x, 70), lbl, font=f, fill=col)
        x += _tw(draw, lbl, f)
        if i < len(parts) - 1:
            draw.text((x, 70), sep, font=f_off, fill=GRAY_L)
            x += _tw(draw, sep, f_off)


def _footer(draw, d: Dict[str, Any], cfg: Dict[str, Any]) -> None:
    fc = cfg.get("factory", {})
    channel = str(fc.get("channel", "rajcwork"))
    f_ch = font(28, "bold")
    ch_w = _tw(draw, channel, f_ch)
    draw.text((W - M - ch_w, 1000), channel, font=f_ch, fill=BLUE)
    left = f"{fc.get('series', 'Software Factory')} · Day {d['day']} · {d['title']}"
    f = font(26)
    while _tw(draw, left, f) > W - 2 * M - ch_w - 60 and len(left) > 20:
        left = left[:-2].rstrip() + "…" if not left.endswith("…") else left[:-3].rstrip() + "…"
    draw.text((M, 1002), left, font=f, fill=GRAY_L)


# --------------------------------------------------------------------------- visuals

def _node_label(n: Any) -> Tuple[str, str]:
    return (n.get("label", ""), n.get("sub", "")) if isinstance(n, dict) else (str(n), "")


def _visual_image(img, draw, v: Dict[str, Any], box) -> None:
    from PIL import Image
    x0, y0, x1, y1 = box
    path = DIAGRAMS_DIR / f"{v['diagram']}.png"
    try:
        src = Image.open(path).convert("RGB")
    except OSError:
        draw.text((x0, y0), f"[diagram {v['diagram']} missing]", font=font(40), fill=GRAY)
        return
    scale = min((x1 - x0) / src.width, (y1 - y0) / src.height)
    size = (max(1, int(src.width * scale)), max(1, int(src.height * scale)))
    src = src.resize(size, Image.LANCZOS)
    img.paste(src, (int(x0 + (x1 - x0 - size[0]) / 2), int(y0 + (y1 - y0 - size[1]) / 2)))


def _visual_flow(draw, v: Dict[str, Any], box) -> None:
    x0, y0, x1, y1 = box
    nodes = [_node_label(n) for n in v["nodes"]]
    n = len(nodes)
    arrows = list(v.get("arrows") or [""] * (n - 1))
    f_arrow = font(28)
    gap = max(130, max((_tw(draw, a, f_arrow) for a in arrows), default=0) + 44)
    bw = (x1 - x0 - gap * (n - 1)) / n
    if bw < 250:
        gap = max(110, (x1 - x0 - 250 * n) / max(1, n - 1))
        bw = (x1 - x0 - gap * (n - 1)) / n
    has_sub = any(sub for _, sub in nodes)
    bh = 300 if has_sub else 250
    loop_h = 180 if v.get("loop") else 0
    top = y0 + ((y1 - y0) - bh - loop_h) / 2
    cy = top + bh / 2
    for i, (label, sub) in enumerate(nodes):
        bx = x0 + i * (bw + gap)
        rounded(draw, (bx, top, bx + bw, top + bh), TINT, BLUE if i == n - 1 else SOFT, 4)
        lab_h = bh * (0.52 if sub else 0.8)
        f, lines, lh = fit(draw, label, bw - 44, lab_h, (56, 52, 48, 44, 40, 36, 32, 28), "bold", max_lines=3)
        sub_lines: List[str] = []
        fs, slh = font(28), 34
        if sub:
            fs, sub_lines, slh = fit(draw, sub, bw - 44, bh * 0.36, (34, 32, 30, 28, 26, 24, 22))
        block = len(lines) * lh + (12 + len(sub_lines) * slh if sub_lines else 0)
        ty = top + (bh - block) / 2
        ty = draw_lines(draw, bx + 22, ty, lines, f, lh, NAVY, "center", bw - 44)
        if sub_lines:
            draw_lines(draw, bx + 22, ty + 12, sub_lines, fs, slh, GRAY, "center", bw - 44)
        if i < n - 1:
            ax0, ax1 = bx + bw + 14, bx + bw + gap - 14
            arrow(draw, (ax0, cy), (ax1, cy))
            if arrows[i]:
                tw = _tw(draw, arrows[i], f_arrow)
                draw.text((ax0 + (ax1 - ax0 - tw) / 2, cy - 54), arrows[i], font=f_arrow, fill=GRAY)
    if v.get("loop"):
        first_cx = x0 + bw / 2
        last_cx = x0 + (n - 1) * (bw + gap) + bw / 2
        ly = top + bh + 80
        draw.line((last_cx, top + bh + 8, last_cx, ly), fill=BLUE, width=5)
        draw.line((last_cx, ly, first_cx, ly), fill=BLUE, width=5)
        arrow(draw, (first_cx, ly), (first_cx, top + bh + 10), BLUE, 5, 20)
        lbl = str(v.get("loop_label", ""))
        if lbl:
            fl, ll, llh = fit(draw, lbl, last_cx - first_cx - 40, 80, (32, 30, 28, 26))
            draw_lines(draw, first_cx + 20, ly + 16, ll, fl, llh, GRAY, "center", last_cx - first_cx - 40)


def _visual_compare(draw, v: Dict[str, Any], box) -> None:
    x0, y0, x1, y1 = box
    mid = 120
    pw = (x1 - x0 - mid) / 2
    for side, px in (("left", x0), ("right", x0 + pw + mid)):
        s = v[side]
        hi = side == "right"
        rounded(draw, (px, y0, px + pw, y1), TINT if hi else NEUTRAL, BLUE if hi else None, 4 if hi else 0)
        ft, tl, tlh = fit(draw, s["title"], pw - 100, 120, (52, 48, 44, 40, 36), "bold", max_lines=2)
        y = draw_lines(draw, px + 50, y0 + 50, tl, ft, tlh, BLUE if hi else GRAY) + 40
        avail = (y1 - 40 - y) / max(1, len(s["points"]))
        for p in s["points"]:
            fp, pl, plh = fit(draw, p, pw - 150, avail - 10, (44, 42, 40, 38, 36, 34, 30, 28))
            dot_y = y + plh / 2 - 8
            draw.ellipse((px + 56, dot_y, px + 74, dot_y + 18), fill=BLUE if hi else GRAY_L)
            draw_lines(draw, px + 98, y, pl, fp, plh, NAVY)
            y += len(pl) * plh + 46
    cx, cy, r = x0 + pw + mid / 2, (y0 + y1) / 2, 48
    draw.ellipse((cx - r, cy - r, cx + r, cy + r), fill=WHITE, outline=SOFT, width=4)
    fv = font(36, "bold")
    draw.text((cx - _tw(draw, "vs", fv) / 2, cy - 24), "vs", font=fv, fill=GRAY)


def _visual_hub(draw, v: Dict[str, Any], box) -> None:
    x0, y0, x1, y1 = box
    cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
    spokes = list(v["spokes"])
    n = len(spokes)
    rx, ry = (x1 - x0) / 2 - 220, (y1 - y0) / 2 - 90
    sw, sh = 340, 120
    offset = 180 / n if n % 2 == 0 else 0
    pos = []
    for i in range(n):
        ang = math.radians(-90 + 360 * i / n + offset)
        pos.append((cx + rx * math.cos(ang), cy + ry * math.sin(ang)))
    for px, py in pos:
        draw.line((cx, cy, px, py), fill=SOFT, width=8)
    cw, ch = 420, 150
    rounded(draw, (cx - cw / 2, cy - ch / 2, cx + cw / 2, cy + ch / 2), BLUE, radius=30)
    f, lines, lh = fit(draw, v["center"], cw - 50, ch - 30, (52, 48, 44, 40, 36), "bold", max_lines=2)
    draw_lines(draw, cx - cw / 2 + 25, cy - len(lines) * lh / 2, lines, f, lh, WHITE, "center", cw - 50)
    for (px, py), label in zip(pos, spokes):
        rounded(draw, (px - sw / 2, py - sh / 2, px + sw / 2, py + sh / 2), TINT, SOFT, 4)
        fs, sl, slh = fit(draw, label, sw - 40, sh - 24, (46, 42, 38, 34, 30), "bold", max_lines=2)
        draw_lines(draw, px - sw / 2 + 20, py - len(sl) * slh / 2, sl, fs, slh, NAVY, "center", sw - 40)


def _visual_layers(draw, v: Dict[str, Any], box) -> None:
    x0, y0, x1, y1 = box
    layers = v["layers"]
    n, gap = len(layers), 20
    bh = min(150, (y1 - y0 - gap * (n - 1)) / n)
    top = y0 + ((y1 - y0) - (bh * n + gap * (n - 1))) / 2
    fl = font(int(min(48, bh * 0.4)), "bold")
    col = max(_tw(draw, x["label"], fl) for x in layers) + 110
    for i, layer in enumerate(layers):
        by = top + i * (bh + gap)
        rounded(draw, (x0, by, x1, by + bh), TINT, SOFT, 3, radius=22)
        draw.rounded_rectangle((int(x0), int(by), int(x0 + 16), int(by + bh)), radius=8, fill=BLUE)
        lh = fl.size if hasattr(fl, "size") else 40
        draw.text((x0 + 56, by + (bh - lh * 1.15) / 2), layer["label"], font=fl, fill=NAVY)
        if layer.get("sub"):
            fs, sl, slh = fit(draw, layer["sub"], x1 - x0 - col - 50, bh - 20, (36, 34, 32, 30, 28, 26, 24))
            draw_lines(draw, x0 + col, by + (bh - len(sl) * slh) / 2, sl, fs, slh, GRAY)


def _visual_code(draw, v: Dict[str, Any], box) -> None:
    x0, y0, x1, y1 = box
    lines = list(v["lines"])
    pad = 54
    size = 36
    for size in (36, 34, 32, 30, 28, 26, 24, 22, 20):
        f = font(size, "mono")
        widest = max((_tw(draw, ln, f) for ln in lines), default=0)
        if widest <= (x1 - x0) - 2 * pad and len(lines) * size * 1.5 <= (y1 - y0) - 2 * pad - 30:
            break
    f = font(size, "mono")
    lh = int(size * 1.5)
    ph = len(lines) * lh + 2 * pad + 30
    py = y0 + ((y1 - y0) - ph) / 2
    rounded(draw, (x0, py, x1, py + ph), NAVY, radius=22)
    lang = str(v.get("lang", "")).upper()
    if lang:
        fl = font(24, "bold")
        draw.text((x1 - pad - _tw(draw, lang, fl), py + 22), lang, font=fl, fill=CODE_DIM)
    y = py + pad + 20
    for ln in lines:
        x = x0 + pad
        if ln.lstrip().startswith("#"):
            draw.text((x, y), ln, font=f, fill=CODE_DIM)
        else:
            body, comment = ln, ""
            cut = ln.find("  #")
            if cut > 0:
                body, comment = ln[:cut], ln[cut:]
            if body.startswith("$ "):
                draw.text((x, y), "$", font=f, fill=CODE_PROMPT)
                x += _tw(draw, "$ ", f)
                body = body[2:]
            draw.text((x, y), body, font=f, fill=CODE_FG)
            if comment:
                draw.text((x + _tw(draw, body, f), y), comment, font=f, fill=CODE_DIM)
        y += lh


def draw_visual(img, draw, v: Dict[str, Any], box) -> None:
    kind = v["type"]
    if kind == "image":
        _visual_image(img, draw, v, box)
    elif kind == "flow":
        _visual_flow(draw, v, box)
    elif kind == "compare":
        _visual_compare(draw, v, box)
    elif kind == "hub":
        _visual_hub(draw, v, box)
    elif kind == "layers":
        _visual_layers(draw, v, box)
    else:
        _visual_code(draw, v, box)


# --------------------------------------------------------------------------- frames

def frame_hook(d: Dict[str, Any], cfg: Dict[str, Any]):
    """What and why in one slide: the concept, one-sentence definition, why it matters."""
    img, draw = _canvas()
    _label(draw, "What & why")
    _tracker(draw, "hook")
    f, lines, lh = fit(draw, d["concept"], W - 2 * M, 230, (120, 112, 104, 96, 88, 80, 72, 64), "bold",
                       spacing=1.08, max_lines=2)
    y = draw_lines(draw, M, 160, lines, f, lh, NAVY) + 26
    fs, sl, slh = fit(draw, d["what"], W - 2 * M, 560 - y, (52, 48, 44, 40, 36, 32), spacing=1.3)
    draw_lines(draw, M, y, sl, fs, slh, NAVY)
    rounded(draw, (M, 600, W - M, CONTENT_BOTTOM - 20), TINT, SOFT, 3)
    draw.text((M + 44, 630), spaced("Why it matters"), font=font(30, "bold"), fill=BLUE)
    fw, wl, wlh = fit(draw, d["why"], W - 2 * M - 88, 170, (46, 42, 40, 38, 36, 34, 30), spacing=1.25)
    draw_lines(draw, M + 44, 690, wl, fw, wlh, NAVY)
    _footer(draw, d, cfg)
    return img


def frame_picture(d: Dict[str, Any], cfg: Dict[str, Any]):
    """Group it: one picture that maps what the demo showed onto the concept."""
    img, draw = _canvas()
    _label(draw, "Group it")
    _tracker(draw, "picture")
    v = d["visual"]
    draw_visual(img, draw, v, (M, 140, W - M, 828))
    caption = str(v.get("caption", ""))
    if v["type"] == "image":
        caption = (caption + "  ·  " if caption else "") + "Diagram: Gas City docs (MIT)"
    if caption:
        fc, cl, clh = fit(draw, caption, W - 2 * M, 50, (32, 30, 28, 26, 24), max_lines=1)
        draw_lines(draw, M, 846, cl, fc, clh, GRAY, "center", W - 2 * M)
    _footer(draw, d, cfg)
    return img


def frame_end(d: Dict[str, Any], cfg: Dict[str, Any]):
    img, draw = _canvas()
    _label(draw, "Takeaway")
    _tracker(draw, "end")
    f, lines, lh = fit(draw, d["takeaway"], W - 2 * M - 80, 560, (92, 86, 80, 74, 68, 62, 56, 50), "bold",
                       spacing=1.2)
    block = len(lines) * lh
    y = 200 + (620 - block) / 2
    draw.rounded_rectangle((M, int(y - 10), M + 16, int(y + block)), radius=8, fill=BLUE)
    draw_lines(draw, M + 70, y, lines, f, lh, NAVY)
    _footer(draw, d, cfg)
    return img


def frame_intro_card(d: Dict[str, Any], cfg: Dict[str, Any], cur: Dict[str, Any]):
    img, draw = _canvas()
    fc = cfg.get("factory", {})
    draw.rectangle((0, 0, 24, H), fill=BLUE)
    draw.text((M, 96), str(fc.get("channel", "rajcwork")), font=font(46, "bold"), fill=BLUE)
    draw.text((M, 330), spaced(f"{fc.get('series', 'Software Factory')} · Day {d['day']}"),
              font=font(36, "bold"), fill=BLUE)
    f, lines, lh = fit(draw, d["title"], W - 2 * M, 300, (112, 104, 96, 88, 80, 72, 64), "bold",
                       spacing=1.1, max_lines=2)
    y = draw_lines(draw, M, 400, lines, f, lh, NAVY) + 24
    idx, n_mods, mod = module_info(cur, d["module"])
    draw.text((M, y), f"Module {idx} of {n_mods} · {mod['title']}", font=font(40), fill=GRAY)
    draw.text((M, 960), "One concept · one takeaway · under three minutes", font=font(32), fill=GRAY_L)
    return img


def frame_outro_card(d: Dict[str, Any], cfg: Dict[str, Any], next_day: Optional[Dict[str, Any]]):
    img, draw = _canvas()
    fc = cfg.get("factory", {})
    f = font(84, "bold")
    msg = "Thanks for watching"
    draw.text(((W - _tw(draw, msg, f)) / 2, 300), msg, font=f, fill=NAVY)
    if next_day:
        k = spaced(f"Tomorrow · Day {next_day['day']}")
        fk = font(34, "bold")
        draw.text(((W - _tw(draw, k, fk)) / 2, 480), k, font=fk, fill=BLUE)
        fn, nl, nlh = fit(draw, next_day["title"], W - 2 * M, 130, (56, 52, 48, 44, 40), max_lines=2)
        draw_lines(draw, M, 540, nl, fn, nlh, GRAY, "center", W - 2 * M)
    else:
        fn, nl, nlh = fit(draw, "That's the whole path. Thank you for learning along.", W - 2 * M, 130,
                          (52, 48, 44))
        draw_lines(draw, M, 500, nl, fn, nlh, GRAY, "center", W - 2 * M)
    ch = str(fc.get("channel", "rajcwork"))
    fch = font(52, "bold")
    draw.text(((W - _tw(draw, ch, fch)) / 2, 840), ch, font=fch, fill=BLUE)
    tag = f"{fc.get('series', 'Software Factory')} · one concept a day"
    ft = font(30)
    draw.text(((W - _tw(draw, tag, ft)) / 2, 915), tag, font=ft, fill=GRAY_L)
    return img


def frame_countdown(d: Dict[str, Any], cfg: Dict[str, Any], take: int):
    img, draw = _canvas()
    max_takes = int(cfg.get("factory", {}).get("max_takes", 3))
    name, text = TAKE_FOCUS.get(take, ("Extra", "Natural and clear. Then publish."))
    draw.text((M, 190), spaced(f"Take {take} of {max_takes} · {name}"), font=font(38, "bold"), fill=BLUE)
    f, lines, lh = fit(draw, text, W - 2 * M, 300, (76, 70, 64, 58, 52, 46), spacing=1.2)
    y = draw_lines(draw, M, 260, lines, f, lh, NAVY) + 40
    draw.text((M, max(y, 620)), f"Day {d['day']} · {d['title']}", font=font(40, "bold"), fill=GRAY)
    secs = segments(cfg)["intro"]
    draw.text((M, max(y, 620) + 70), f"Starting in {secs} seconds. Breathe. When the bar ends: what it is, and why.",
              font=font(34), fill=GRAY_L)
    _footer(draw, d, cfg)
    return img


def frame_closing(d: Dict[str, Any], cfg: Dict[str, Any]):
    img, draw = _canvas()
    draw.text((M, 190), spaced("Take done"), font=font(38, "bold"), fill=BLUE)
    draw.text((M, 260), "Recording stops on its own.", font=font(84, "bold"), fill=NAVY)
    draw.text((M, 400), "Next: listen once, then choose — another take, or publish.", font=font(44), fill=GRAY)
    _footer(draw, d, cfg)
    return img


def render_frames(d: Dict[str, Any], cfg: Dict[str, Any], cur: Dict[str, Any],
                  next_day: Optional[Dict[str, Any]], out_dir: Path) -> Dict[str, Path]:
    """Render every content frame of a day. Returns {name: png path}."""
    out_dir.mkdir(parents=True, exist_ok=True)
    resolve_fonts(cfg)
    frames = {
        "intro": frame_intro_card(d, cfg, cur),
        "hook": frame_hook(d, cfg),
        "picture": frame_picture(d, cfg),
        "end": frame_end(d, cfg),
        "outro": frame_outro_card(d, cfg, next_day),
        "closing": frame_closing(d, cfg),
    }
    paths: Dict[str, Path] = {}
    for name, im in frames.items():
        p = out_dir / f"{name}.png"
        im.save(p, "PNG")
        paths[name] = p
    return paths


def render_countdown(d: Dict[str, Any], cfg: Dict[str, Any], take: int, out_dir: Path) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    p = out_dir / f"countdown_take{take}.png"
    frame_countdown(d, cfg, take).save(p, "PNG")
    return p


def render_thumbnail(d: Dict[str, Any], cfg: Dict[str, Any], out_png: Path) -> Path:
    """1280x720 YouTube thumbnail in the same look as the intro card."""
    from PIL import Image, ImageDraw
    resolve_fonts(cfg)
    fc = cfg.get("factory", {})
    img = Image.new("RGB", (1280, 720), WHITE)
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, 18, 720), fill=BLUE)
    draw.text((70, 64), spaced(f"{fc.get('series', 'Software Factory')} · Day {d['day']}"),
              font=font(30, "bold"), fill=BLUE)
    f, lines, lh = fit(draw, d["title"], 1140, 380, (96, 88, 80, 72, 64, 58, 52), "bold", spacing=1.1,
                       max_lines=3)
    draw_lines(draw, 70, 140, lines, f, lh, NAVY)
    draw.rounded_rectangle((70, 590, 70 + 120, 600), radius=5, fill=BLUE)
    draw.text((70, 618), str(fc.get("channel", "rajcwork")), font=font(40, "bold"), fill=BLUE)
    tag = "One concept a day"
    ft = font(32)
    draw.text((1280 - 70 - _tw(draw, tag, ft), 624), tag, font=ft, fill=GRAY_L)
    out_png.parent.mkdir(parents=True, exist_ok=True)
    img.save(out_png, "PNG")
    return out_png


# --------------------------------------------------------------------------- demo (animated terminal)

PROMPTS = {"rig": "hello-factory $", "city": "lab $", "scratch": "scratch $"}


def demo_steps_for(d: Dict[str, Any], cap: Optional[Dict[str, Any]]) -> Tuple[List[Dict[str, Any]], bool]:
    """Steps to show: the real capture when there is a good one, else a preview of the commands."""
    if cap and cap.get("ok") and cap.get("topic") == d["id"] and cap.get("steps"):
        return list(cap["steps"]), False
    from .factory import demo_steps
    shown = [st for st in d["demo"]["steps"] if "write" not in st and not st.get("hidden")]
    steps = []
    for (cmd, say), st in zip(demo_steps(d), shown):
        steps.append({"cmd": cmd.split("   # time-lapse")[0], "say": say,
                      "output": list(st.get("expect") or ["(the real output appears after: fde-coach factory demo capture)"]),
                      "highlight": "", "timelapse": 60 if "wait_for" in st else 0, "cwd": st.get("cwd", "rig")})
    return steps, True


def allocate(steps: Sequence[Dict[str, Any]], total: float) -> List[float]:
    """Split the demo time across steps: more output or a time-lapse gets more time; 6 s minimum."""
    weights = [2 + min(len(st.get("output") or []), 14) / 3 + (1.5 if st.get("timelapse") else 0) for st in steps]
    raw = [max(6.0, total * w / sum(weights)) for w in weights]
    scale = total / sum(raw)
    out = [round(x * scale, 1) for x in raw]
    out[-1] = round(total - sum(out[:-1]), 1)
    return out


def _lapse(seconds: int) -> str:
    m, s_ = divmod(max(1, int(seconds)), 60)
    return f"» {m} min {s_} s later" if m else f"» {s_} s later"


def frame_demo(d: Dict[str, Any], cfg: Dict[str, Any], lines: List[Tuple[str, str]], idx: int, n: int,
               say: str, highlight: str, preview: bool):
    """lines: [(kind, text)] with kind prompt | out | lapse; the newest lines are shown."""
    img, draw = _canvas()
    _label(draw, "See it")
    _tracker(draw, "demo")
    ft = font(46, "bold")
    draw.text((M, 128), d["demo"]["title"], font=ft, fill=NAVY)
    fs = font(26)
    step = f"Step {idx} of {n}"
    draw.text((W - M - _tw(draw, step, fs), 142), step, font=fs, fill=GRAY_L)
    x0, y0, x1, y1 = M, 208, W - M, 840
    rounded(draw, (x0, y0, x1, y1), TERM_BG, radius=18)
    draw.rounded_rectangle((x0, y0, x1, y0 + 46), radius=18, fill=TERM_BAR)
    draw.rectangle((x0, y0 + 30, x1, y0 + 46), fill=TERM_BAR)
    for i, col in enumerate(("#E5604F", "#F2B84B", "#5CC26A")):
        draw.ellipse((x0 + 22 + i * 30, y0 + 15, x0 + 38 + i * 30, y0 + 31), fill=col)
    path = "~/hello-factory" if d["demo"].get("store") else "~"
    store = "file store" if d["demo"].get("store") == "file" else "bd + Dolt"
    fh = font(22)
    head = f"{path}  ·  rajcwork demo city ({store})"
    draw.text(((x0 + x1 - _tw(draw, head, fh)) / 2, y0 + 11), head, font=fh, fill=CODE_DIM)
    if preview:
        fp = font(22, "bold")
        tag = "PREVIEW — not captured yet"
        draw.text((x1 - 24 - _tw(draw, tag, fp), y0 + 11), tag, font=fp, fill=AMBER)
    fm = font(26, "mono")
    lh = 37
    cw = max(1.0, _tw(draw, "0" * 10, fm) / 10)
    cols = int((x1 - x0 - 64) / cw)
    rows: List[Tuple[str, str]] = []
    for kind, text in lines:
        if kind == "lapse":
            rows.append((kind, text))
            continue
        chunks = [text[i:i + cols] for i in range(0, max(1, len(text)), cols)] or [""]
        rows += [(kind if j == 0 else ("out" if kind == "out" else "cont"), c) for j, c in enumerate(chunks)]
    fit_rows = int((y1 - y0 - 80) / lh)
    rows = rows[-fit_rows:]
    y = y0 + 64
    hi = highlight.lower().strip()
    for kind, text in rows:
        if kind == "lapse":
            fl = font(28, "bold")
            w = _tw(draw, text, fl) + 48
            cx = (x0 + x1) / 2
            rounded(draw, (cx - w / 2, y - 2, cx + w / 2, y + lh - 4), BLUE, radius=16)
            draw.text((cx - w / 2 + 24, y + 2), text, font=fl, fill=WHITE)
        elif kind in ("prompt",):
            prompt, _, cmd = text.partition("\u0000")
            draw.text((x0 + 32, y), prompt, font=fm, fill=CODE_PROMPT)
            draw.text((x0 + 32 + _tw(draw, prompt + " ", fm), y), cmd, font=fm, fill=WHITE)
        else:
            if hi and hi in text.lower():
                draw.rectangle((x0 + 20, y - 3, x1 - 20, y + lh - 5), fill=TERM_HI)
                draw.text((x0 + 32, y), text, font=fm, fill=WHITE)
            else:
                draw.text((x0 + 32, y), text, font=fm, fill=CODE_FG if kind == "cont" else TERM_OUT)
        y += lh
    if say:
        fc, cl, clh = fit(draw, say, W - 2 * M, 50, (32, 30, 28, 26), max_lines=1)
        draw_lines(draw, M, 856, cl, fc, clh, GRAY, "center", W - 2 * M)
    _footer(draw, d, cfg)
    return img


def render_demo(d: Dict[str, Any], cfg: Dict[str, Any], cap: Optional[Dict[str, Any]],
                out_dir: Path) -> Dict[str, Any]:
    """Animated terminal for the demo section. Returns {timeline: [(png, s)], slides: [(png, s, notes)],
    captured: bool}. Commands type out, real output appears, slow agent work becomes a time-lapse card."""
    out_dir.mkdir(parents=True, exist_ok=True)
    for old in out_dir.glob("*.png"):
        old.unlink()
    steps, preview = demo_steps_for(d, cap)
    secs = allocate(steps, float(segments(cfg)["demo"]))
    history: List[Tuple[str, str]] = []
    timeline: List[Tuple[Path, float]] = []
    slides: List[Tuple[Path, float, str]] = []
    k = 0

    def save(im, dur: float) -> Path:
        nonlocal k
        k += 1
        pth = out_dir / f"demo_{k:03d}.png"
        im.save(pth, "PNG")
        timeline.append((pth, round(dur, 2)))
        return pth

    n = len(steps)
    for i, (st, t) in enumerate(zip(steps, secs), 1):
        prompt = PROMPTS.get(st.get("cwd", "rig"), "$")
        cmd = st["cmd"]
        n_type = min(6, max(2, len(cmd) // 14))
        typing = 0.22
        for j in range(1, n_type + 1):
            part = cmd[: int(len(cmd) * j / n_type)]
            cursor = "█" if j < n_type else ""
            save(frame_demo(d, cfg, history + [("prompt", f"{prompt}\u0000{part}{cursor}")], i, n, st.get("say", ""),
                            "", preview), typing)
        history.append(("prompt", f"{prompt}\u0000{cmd}"))
        used = n_type * typing
        if st.get("timelapse"):
            save(frame_demo(d, cfg, history + [("lapse", _lapse(st["timelapse"]))], i, n, st.get("say", ""), "",
                            preview), 1.6)
            history.append(("lapse", _lapse(st["timelapse"])))
            used += 1.6
        history += [("out", ln) for ln in (st.get("output") or [])]
        final = frame_demo(d, cfg, history, i, n, st.get("say", ""), st.get("highlight", ""), preview)
        pth = save(final, max(0.5, t - used))
        notes = f"Step {i} of {n}: {cmd}\n\nPoint out: {st.get('say', '')}"
        if st.get("highlight"):
            notes += f"\nLook at: {st['highlight']}"
        slides.append((pth, t, notes))
    return {"timeline": timeline, "slides": slides, "captured": not preview}


# --------------------------------------------------------------------------- take deck (.pptx)

def _speaker_notes(d: Dict[str, Any], section: str) -> str:
    if section == "hook":
        lines = [f"Open with: Today: {d['concept'].lower()}.", d["what"], "", f"Why: {d['why']}", ""]
        lines += [f"Explain '{j['term']}' right away: {j['plain']}" for j in d["jargon"]]
        lines += ["", f"Then: let's see it. ({d['demo']['title']})"]
    elif section == "picture":
        lines = ["Group what they just saw:", d["how"], ""] + [f"- {p}" for p in d["how_points"]]
        lines += ["", f"Analogy: {d['analogy']}"]
    else:
        lines = [f"Close with: {d['takeaway']}", "", "Then stop talking; the outro music plays in the video."]
    return "\n".join(lines)


def build_take_deck(d: Dict[str, Any], cfg: Dict[str, Any], frames: Dict[str, Path], countdown: Path,
                    demo_slides: Sequence[Tuple[Any, float, str]], out_dir: Path, stem: str) -> Tuple[Path, Path]:
    """Countdown, hook, one slide per demo step, picture, takeaway (all auto-advancing with a progress
    bar), then a closing slide."""
    from pptx import Presentation
    from pptx.util import Emu, Inches

    from .deck import (SLIDE_H, SLIDE_W, _ensure_use_timings, _notes, _register_notes_master, _timer,
                       make_show_copy, set_auto_animations, set_transition)

    out_dir.mkdir(parents=True, exist_ok=True)
    seg = segments(cfg)
    prs = Presentation()
    prs.slide_width = Emu(int(Inches(SLIDE_W)))
    prs.slide_height = Emu(int(Inches(SLIDE_H)))
    plan = [(countdown, seg["intro"], "Take is recording. Settle; start when the slide changes.")]
    plan.append((frames["hook"], seg["hook"], _speaker_notes(d, "hook")))
    plan += [(png, secs, notes) for png, secs, notes in demo_slides]
    plan.append((frames["picture"], seg["picture"], _speaker_notes(d, "picture")))
    plan.append((frames["end"], seg["end"], _speaker_notes(d, "end")))
    for png, secs, notes in plan:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_picture(str(png), 0, 0, width=prs.slide_width, height=prs.slide_height)
        effects = _timer(s, secs, ticks=False)
        set_transition(s, int(round(secs * 1000)))
        set_auto_animations(s, effects)
        _notes(s, notes)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(str(frames["closing"]), 0, 0, width=prs.slide_width, height=prs.slide_height)
    set_transition(s, None)
    _register_notes_master(prs)
    _ensure_use_timings(prs)
    cp = prs.core_properties
    cp.title = f"{cfg.get('factory', {}).get('series', 'Software Factory')} — Day {d['day']}: {d['title']}"
    cp.author = "FDE Impromptu Coach"
    pptx_path = out_dir / f"{stem}.pptx"
    tmp = out_dir / f".{stem}.tmp.pptx"
    prs.save(str(tmp))
    tmp.replace(pptx_path)
    ppsx_path = out_dir / f"{stem}.ppsx"
    make_show_copy(pptx_path, ppsx_path)
    return pptx_path, ppsx_path


def take_seconds(cfg: Dict[str, Any]) -> int:
    seg = segments(cfg)
    return sum(seg[s] for s in ("intro", "hook", "demo", "picture", "end"))
