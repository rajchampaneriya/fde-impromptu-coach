"""Software Factory video series tests: curriculum integrity, day planning, the
briefs, slide rendering, the dry-run brief -> kit -> takes -> publish flow, and
a real ffmpeg assembly check (skipped without ffmpeg)."""
from __future__ import annotations

import datetime as dt
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

SCRIPTS = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SCRIPTS))

from fdecoach import audio, factory as F, gcal  # noqa: E402
from fdecoach.config import DEFAULTS, REFERENCES_DIR, Paths, load_config  # noqa: E402
from fdecoach.state import History  # noqa: E402

from test_coach import Home  # noqa: E402  (same dry-run harness)

START = dt.date(2026, 10, 10)


def _history(tmp: Path) -> History:
    paths = Paths(tmp)
    paths.ensure()
    return History(paths, file=paths.factory_state)


class TestCurriculum(unittest.TestCase):
    def test_curriculum_is_valid(self):
        cur = F.load_curriculum()
        self.assertEqual(F.validate_curriculum(cur), [])
        self.assertEqual(F.total_days(cur), 35)
        used = {d["module"] for d in cur["days"]}
        self.assertEqual(used, {m["id"] for m in cur["modules"]}, "every module has days")

    def test_one_concept_one_takeaway(self):
        for d in F.load_curriculum()["days"]:
            self.assertEqual(d["takeaway"].count(". "), 0, f"day {d['day']}: takeaway is one sentence")
            self.assertTrue(d["takeaway"].endswith("."), d["day"])

    def test_validation_catches_problems(self):
        d = json.loads(json.dumps(F.load_curriculum()["days"][0]))
        d["what_points"] = ["x" * 80]
        d["visual"] = {"type": "flow", "nodes": ["only one"]}
        d["references"] = ["https://example.com/doc"]
        errs = " | ".join(F.validate_day(d))
        self.assertIn("what_points", errs)
        self.assertIn("2-5 nodes", errs)
        self.assertIn("relative to the repository root", errs)

    def test_video_under_three_minutes_with_valid_chapters(self):
        self.assertLessEqual(F.video_seconds(DEFAULTS), F.MAX_VIDEO_SECONDS)
        marks = F.chapters(F.topic(1), DEFAULTS)
        self.assertEqual(marks[0][0], 0)
        ends = [t for t, _ in marks[1:]] + [F.video_seconds(DEFAULTS)]
        for (t, _), end in zip(marks, ends):
            self.assertGreaterEqual(end - t, 10, "YouTube chapters must be at least 10 s")

    def test_learning_path_doc_in_sync(self):
        cfg = load_config(Paths(Path(tempfile.mkdtemp())))
        want = F.learning_path_doc(cfg).markdown()
        have = (REFERENCES_DIR / "factory_learning_path.md").read_text(encoding="utf-8")
        self.assertEqual(have, want, "regenerate: fde-coach factory plan --markdown --write "
                                     "references/factory_learning_path.md")


class TestPlanning(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.h = _history(self.tmp)
        self.cfg = DEFAULTS

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def put(self, date: str, day: int, recorded: bool) -> None:
        self.h.put({"date": date, "day": day, "recorded": recorded, "takes": []})

    def test_before_and_on_start(self):
        self.assertIsNone(F.day_for(self.h, self.cfg, START - dt.timedelta(days=1), START))
        self.assertEqual(F.day_for(self.h, self.cfg, START, START - dt.timedelta(days=1)), 1)
        self.assertEqual(F.day_for(self.h, self.cfg, START + dt.timedelta(days=4), dt.date(2026, 10, 1)), 5)

    def test_missed_day_carries_over(self):
        self.put("2026-10-10", 1, recorded=False)
        self.assertEqual(F.day_for(self.h, self.cfg, dt.date(2026, 10, 11), dt.date(2026, 10, 11)), 1)

    def test_published_moves_on_and_projects_forward(self):
        self.put("2026-10-10", 1, recorded=True)
        ref = dt.date(2026, 10, 11)
        self.assertEqual(F.day_for(self.h, self.cfg, ref, ref), 2)
        self.assertEqual(F.day_for(self.h, self.cfg, dt.date(2026, 10, 14), ref), 5)

    def test_beyond_the_path(self):
        self.assertIsNone(F.topic(F.total_days() + 1))
        self.assertIsNone(F.topic(0))


class TestDocuments(unittest.TestCase):
    def test_brief_has_every_part(self):
        d = F.topic(3)
        md = F.brief_doc(d, START + dt.timedelta(days=2), DEFAULTS).markdown()
        for part in ("Tomorrow: The Six Primitives", "## WHAT", "## WHY", "## HOW",
                     "## Key repository references", "## Relevant existing diagrams",
                     "## Artifact or demo to prepare", "## 3-minute outline", "Monday 12 October 2026",
                     "docs/diagrams/excalidraw-rendered/primitives.svg"):
            self.assertIn(part, md)

    def test_prep_sheet_and_html(self):
        d = F.topic(7)
        doc = F.prep_doc(d, START + dt.timedelta(days=6), DEFAULTS, {"deck": "/x/deck.pptx"})
        md, page = doc.markdown(), doc.html()
        for part in ("Speaking outline", "## Slides", "## Diagram", "## Demo or artifact",
                     "Key technical points to verify", "Final 3-minute flow", "Take 3 — Publish"):
            self.assertIn(part, md)
        self.assertIn("- [ ] Tutorial 01", md)
        self.assertTrue(page.startswith("<!doctype html>"))
        self.assertIn("<pre><code>gc init ~/my-city\ngc rig add ~/my-project", page)
        self.assertNotIn("<script", page)

    def test_override_is_merged_and_invalid_ignored(self):
        tmp = Path(tempfile.mkdtemp())
        try:
            paths = Paths(tmp)
            path = F.override_path(paths, "software-factory")
            path.parent.mkdir(parents=True)
            path.write_text(json.dumps({"takeaway": "Factories finish work; agents start it.", "day": 99}))
            d = F.topic(1, paths)
            self.assertEqual(d["takeaway"], "Factories finish work; agents start it.")
            self.assertEqual(d["day"], 1, "day can't be overridden")
            path.write_text(json.dumps({"what_points": ["way too long " * 10]}))
            self.assertEqual(F.topic(1, paths)["what_points"], F.topic(1)["what_points"])
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


class TestRendering(unittest.TestCase):
    def test_every_day_renders(self):
        from fdecoach import factory_deck as FD
        for d in F.load_curriculum()["days"]:
            for fn in (FD.frame_what, FD.frame_why, FD.frame_how, FD.frame_end):
                img = fn(d, DEFAULTS)
                self.assertEqual(img.size, (1920, 1080), f"day {d['day']} {fn.__name__}")

    def test_thumbnail(self):
        from fdecoach import factory_deck as FD
        tmp = Path(tempfile.mkdtemp())
        try:
            from PIL import Image
            p = FD.render_thumbnail(F.topic(35), DEFAULTS, tmp / "t.png")
            with Image.open(p) as im:
                self.assertEqual(im.size, (1280, 720))
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


class TestEventIds(unittest.TestCase):
    def test_factory_prefixes(self):
        from fdecoach import factory_session as FS
        for prefix in (FS.GCAL_BRIEF, FS.GCAL_PUBLISH):
            eid = gcal.event_id(dt.date(2026, 10, 10), prefix)
            self.assertRegex(eid, r"^[a-v0-9]{5,}$")
            self.assertEqual(FS._eid_date(eid), "2026-10-10")


class TestFactoryEndToEnd(unittest.TestCase):
    def setUp(self):
        self.h = Home()

    def tearDown(self):
        self.h.cleanup()

    def state(self) -> dict:
        path = self.h.dir / "home" / "state" / "factory.json"
        return json.loads(path.read_text()) if path.exists() else {"sessions": {}}

    def calendar(self) -> dict:
        path = self.h.dir / "home" / "state" / "dryrun_calendar.json"
        return json.loads(path.read_text()) if path.exists() else {}

    def test_brief_kit_takes_publish(self):
        self.assertTrue((self.h.dir / "home" / "Start Video.command").exists())
        # the evening before the start: the brief arrives and the calendar fills
        out = self.h.run("remind", now="2026-10-09T19:05:00", answers="Later").stderr
        self.assertIn("Tomorrow's video | Day 1: Software Factory vs. Coding Agent", out)
        briefs = list((self.h.dir / "home" / "factory" / "briefs").glob("2026-10-10_Day01_*.html"))
        self.assertEqual(len(briefs), 1)
        cal = self.calendar()
        self.assertEqual(cal["fdefpub20261010"]["status"], "confirmed")
        self.assertEqual(cal["fdefbrief20261011"]["status"], "confirmed")
        body = cal["fdefbrief20261011"]["body"]
        self.assertIn("T19:00:00", body["start"]["dateTime"])
        self.assertTrue(body["start"]["dateTime"].startswith("2026-10-10"))
        self.assertIn("Gas City: Zero Hardcoded Roles", body["summary"])
        self.assertIn("WHAT:", body["description"])
        # a second tick doesn't repeat the brief
        again = self.h.run("remind", now="2026-10-09T19:20:00", answers="Later").stderr
        self.assertNotIn("Tomorrow's video", again)

        # recording morning: the kit is built and opened
        out = self.h.run("remind", now="2026-10-10T07:15:00", answers="Later").stderr
        self.assertIn("Today's video kit is ready", out)
        s = self.state()["sessions"]["2026-10-10"]
        self.assertEqual(s["day"], 1)
        self.assertTrue(Path(s["kit"]["deck"]).exists())
        self.assertTrue(Path(s["kit"]["prep_html"]).exists())

        # take 1, publish it, paste the link
        res = self.h.run("factory", "take", now="2026-10-10T08:00:00",
                         answers="Start take 1,Publish this take,https://youtu.be/abc123")
        self.assertIn('"url": "https://youtu.be/abc123"', res.stdout)
        s = self.state()["sessions"]["2026-10-10"]
        self.assertTrue(s["recorded"])
        self.assertEqual(s["final_take"], 1)
        self.assertEqual(len(s["takes"]), 1)
        self.assertEqual(s["youtube"]["status"], "published")
        folder = Path(s["video_path"]).parent
        for name in ("final.mp4", "thumbnail.png", "youtube.txt", "take1.m4a"):
            self.assertTrue((folder / name).exists(), name)
        self.assertEqual(self.calendar()["fdefpub20261010"]["status"], "cancelled")
        st = json.loads(self.h.run("status", "--json", now="2026-10-10T09:00:00").stdout)
        self.assertEqual(st["factory"]["streak"]["current"], 1)
        self.assertEqual(st["factory"]["tomorrow_topic"]["day"], 2)
        self.assertEqual(st["streak"]["current"], 0, "the FDE streak is separate")

    def test_three_takes_then_choose(self):
        self.h.run("factory", "take", now="2026-10-10T08:00:00",
                   answers="Start take 1,Start take 2,Start take 3,Later")
        s = self.state()["sessions"]["2026-10-10"]
        self.assertEqual([t["n"] for t in s["takes"]], [1, 2, 3])
        self.assertFalse(s["recorded"])
        out = self.h.run("factory", "take", now="2026-10-10T09:00:00",
                         answers="Choose a take,Take 2,https://youtu.be/xyz")
        self.assertIn("https://youtu.be/xyz", out.stdout)
        s = self.state()["sessions"]["2026-10-10"]
        self.assertEqual(len(s["takes"]), 3, "no fourth take")
        self.assertEqual(s["final_take"], 2)

    def test_missed_day_carries_over_and_published_later(self):
        self.h.run("factory", "prep", now="2026-10-10T09:00:00")
        out = self.h.run("factory", "status", now="2026-10-11T09:00:00").stdout
        self.assertIn("Today: Day 1 · Software Factory vs. Coding Agent", out)
        self.h.run("factory", "take", now="2026-10-11T09:05:00", answers="Start take 1,Publish this take,")
        s = self.state()["sessions"]["2026-10-11"]
        self.assertTrue(s["carried_over"])
        self.assertFalse(s["recorded"], "no link yet")
        self.assertEqual(s["youtube"]["status"], "ready")
        self.h.run("factory", "published", "--url", "https://youtu.be/late1", now="2026-10-11T10:00:00")
        rows = json.loads(self.h.run("factory", "plan", "--json", now="2026-10-11T10:01:00").stdout)
        self.assertEqual(rows[0]["status"], "published")
        self.assertEqual(rows[1]["date"], "2026-10-12")

    def test_api_mode_uploads(self):
        self.h.run("config", "--set", "factory.publish_mode=api", now="2026-10-10T07:00:00")
        res = self.h.run("factory", "take", now="2026-10-10T08:00:00", answers="Start take 1,Publish this take")
        self.assertIn('"mode": "api"', res.stdout)
        s = self.state()["sessions"]["2026-10-10"]
        self.assertTrue(s["youtube"]["url"].startswith("https://youtu.be/DRYRUN"))

    def test_take_deck_structure(self):
        out = json.loads(self.h.run("factory", "prep", now="2026-10-10T07:00:00").stdout)
        from pptx import Presentation
        prs = Presentation(out["deck"])
        self.assertEqual(len(prs.slides), 6)
        with zipfile.ZipFile(out["deck"]) as z:
            for i, adv in enumerate(("5000", "35000", "35000", "75000", "15000"), 1):
                self.assertIn(f'advTm="{adv}"', z.read(f"ppt/slides/slide{i}.xml").decode(), i)
            self.assertNotIn("advTm", z.read("ppt/slides/slide6.xml").decode())
            self.assertIn("notesMasterIdLst", z.read("ppt/presentation.xml").decode())
        self.assertIn("Open with: Today:", prs.slides[1].notes_slide.notes_text_frame.text)

    def test_before_start_and_brief_cli(self):
        res = self.h.run("factory", "take", now="2026-10-05T08:00:00", check=False)
        self.assertNotEqual(res.returncode, 0)
        self.assertIn("starts on 2026-10-10", res.stderr)
        out = self.h.run("factory", "brief", now="2026-10-09T10:00:00").stdout
        self.assertIn("# Tomorrow: Software Factory vs. Coding Agent", out)
        out = self.h.run("factory", "brief", "--date", "2026-10-12", now="2026-10-09T10:00:00").stdout
        self.assertIn("Day 3", out)

    def test_metadata(self):
        from fdecoach import factory_session as FS
        cfg = load_config(Paths(self.h.dir / "unused"))
        meta = FS.youtube_metadata(F.topic(12), cfg)
        sn = meta["snippet"]
        self.assertLessEqual(len(sn["title"]), 100)
        self.assertIn("0:00 What: Convoy", sn["description"])
        self.assertIn("1:15 How it works", sn["description"])
        self.assertIn("Next: Day 13", sn["description"])
        self.assertEqual(len(sn["tags"]), len(set(sn["tags"])))
        self.assertEqual(meta["status"]["privacyStatus"], "public")


class TestFfmpegAssembly(unittest.TestCase):
    """Real ffmpeg with short sections: the video is exactly as long as the plan, and the voice
    keeps its recorded level (no processing, no mono-to-stereo loss)."""

    def test_assembly_length_and_voice_level(self):
        if not audio.has_ffmpeg() or not shutil.which("ffprobe"):
            self.skipTest("ffmpeg not installed")
        tmp = Path(tempfile.mkdtemp())
        try:
            home = tmp / "home"
            home.mkdir()
            short = {f"{k}_seconds": v for k, v in
                     (("intro", 1), ("what", 3), ("why", 2), ("how", 3), ("end", 2), ("outro", 2))}
            (home / "config.json").write_text(json.dumps({"factory": short}))
            exe = audio.ffmpeg_path()
            voice = tmp / "voice.m4a"
            subprocess.run([exe, "-y", "-f", "lavfi", "-i", "sine=frequency=330:duration=14,volume=0.5",
                            "-ac", "1", "-c:a", "aac", "-b:a", "128k", str(voice)],
                           capture_output=True, check=True, timeout=60)
            env = {k: v for k, v in os.environ.items() if k not in ("FDE_COACH_DRYRUN", "FDE_COACH_RECORDINGS")}
            env.update(FDE_COACH_HOME=str(home), FDE_COACH_RECORDINGS=str(tmp / "movies"),
                       FDE_COACH_NOW="2026-10-10T09:00:00")
            with mock.patch.dict(os.environ, env, clear=True):
                from fdecoach import factory_session as FS
                from fdecoach.app import Ctx
                ctx = Ctx()
                s = FS.ensure_kit(ctx)
                video = FS.assemble_video(ctx, s, F.topic(1), {"n": 1, "audio": str(voice), "offset": 1.0})
            self.assertIsNotNone(video)
            dur = audio.audio_duration_seconds(Path(video))
            self.assertAlmostEqual(dur, 13.0, delta=0.3)

            def mean_db(path: Path, start: float, end: float) -> float:
                err = subprocess.run([exe, "-hide_banner", "-ss", str(start), "-to", str(end), "-i", str(path),
                                      "-vn", "-af", "volumedetect", "-f", "null", "-"],
                                     capture_output=True, text=True, timeout=60).stderr
                return float(re.search(r"mean_volume: (-?[0-9.]+) dB", err).group(1))

            self.assertAlmostEqual(mean_db(Path(video), 4, 9), mean_db(voice, 5, 10), delta=0.6)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
