"""Behaviour tests for exec mode and the exec answer key.

The scenarios live in tests/features/*.feature (Gherkin, the readable spec).
Each scenario has exactly one test here, named after it, written Given / When /
Then; TestFeatureFilesStayInSync fails if a scenario has no test or a test no
longer has a scenario. Run from the skill folder:

    python3 -m unittest discover -s scripts/tests -v
"""
from __future__ import annotations

import datetime as dt
import json
import re
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path

from test_coach import Home  # noqa: E402  (also puts scripts/ on sys.path)

from fdecoach import answers as A  # noqa: E402
from fdecoach import questions as Q  # noqa: E402
from fdecoach.config import DEFAULTS, Paths, load_config  # noqa: E402
from fdecoach.deck import build_deck, chapters, format_line, session_seconds  # noqa: E402

FEATURES = Path(__file__).resolve().parent / "features"
SENIOR = ("exec", "account_leader")
SHAPE = "Headline → So what → What I need → Offer detail"


def scenarios(feature: str):
    text = (FEATURES / feature).read_text(encoding="utf-8")
    return re.findall(r"^\s*Scenario: (.+?)\s*$", text, re.MULTILINE)


def scenario_test_name(scenario: str) -> str:
    slug = re.sub(r"['’]", "", scenario.lower())
    return "test_" + re.sub(r"[^a-z0-9]+", "_", slug).strip("_")


def slides(deck_path: str):
    """[(slide text, notes text)] in show order."""
    from pptx import Presentation
    out = []
    for s in Presentation(deck_path).slides:
        text = "\n".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        out.append((text, notes))
    return out


def bank_question(qid: str) -> dict:
    return dict(next(q for q in Q.bank()["questions"] if q["id"] == qid))


def finished(qids, cfg, date=dt.date(2026, 10, 2)):
    qs = []
    for i, qid in enumerate(qids, 1):
        q = Q._finish(bank_question(qid), "bank")
        q["slot"] = i
        qs.append(q)
    return Q.apply_exec_mode(qs, cfg, date)


def with_reply(q: dict, kind: str) -> dict:
    r = Q.bank()["exec_replies"][kind]
    q["exec_reply"] = {"kind": kind, "line": r["lines"][q["audience"]], "seconds": r["seconds"],
                       "instruction": r["instruction"], "check": r["check"]}
    return q


def cards(page: Path):
    return page.read_text(encoding="utf-8").split('<section class="card"')[1:]


# --------------------------------------------------------------------------- Feature: exec slots

class ExecSlotsFeature(unittest.TestCase):
    feature = "exec_mode.feature"

    def setUp(self):
        self.h = None
        self.tmp = Path(tempfile.mkdtemp(prefix="fdecoach-exec-"))
        self.cfg = load_config(Paths(self.tmp / "unused"))

    def tearDown(self):
        if self.h:
            self.h.cleanup()
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_exec_focus_fills_slots_with_senior_listeners(self):
        # Given focus_audiences is ["exec", "account_leader"]
        self.h = Home()
        self.h.run("config", "--set", 'exec_mode.focus_audiences=["exec", "account_leader"]',
                   now="2026-10-01T05:00:00")
        # When today's set is planned (and built)
        self.h.run("generate", now="2026-10-01T05:30:00")
        s = self.h.history()["sessions"]["2026-10-01"]
        # Then at least 2 of 5 slots have an exec or account-leader listener
        senior = [q for q in s["questions"] if q["audience"] in SENIOR]
        self.assertEqual(s["source"], "claude")
        self.assertGreaterEqual(len(senior), 2)
        self.assertEqual({q["audience"] for q in senior} & set(SENIOR), set(SENIOR), "one of each")
        # And each of those slides shows the listener and what they care about
        shown = slides(s["deck_path"])
        for q in senior:
            text = next(t for t, _ in shown if q["text"] in t)
            self.assertIn(q["listener"], text)
            self.assertIn("cares about " + q["cares"], text)

    def test_exec_focus_survives_the_offline_fallback(self):
        # Given Claude is unreachable
        self.h = Home(claude=False)
        # When today's set is built from the question bank
        self.h.run("generate", now="2026-10-01T05:30:00")
        s = self.h.history()["sessions"]["2026-10-01"]
        self.assertEqual(s["source"], "bank")
        # Then at least 2 of 5 questions still have an exec or account-leader listener
        self.assertGreaterEqual(sum(q["audience"] in SENIOR for q in s["questions"]), 2)
        self.assertTrue(all(q["exec"] for q in s["questions"] if q["audience"] in SENIOR))

    def test_a_question_written_for_the_wrong_listener_is_rejected(self):
        # Given a slot that needs an exec or account-leader listener
        self.h = Home()
        # When Claude writes that question for a technical peer
        self.h.run("generate", now="2026-10-01T05:30:00", extra={"FAKE_CLAUDE_MODE": "wrong_listener"})
        s = self.h.history()["sessions"]["2026-10-01"]
        # Then the question is rejected with the reason
        self.assertTrue(any("the listener must be audience" in n for n in s["notes"]), s["notes"])
        # And the slot is filled from the bank with the listener it needs
        from_bank = [q for q in s["questions"] if q["source"] == "bank"]
        self.assertGreaterEqual(len(from_bank), 2)
        self.assertGreaterEqual(sum(q["audience"] in SENIOR for q in from_bank), 2)

    def test_exec_slides_use_the_headline_window(self):
        # Given a slot whose listener is "exec"
        qs = finished(["EXE-04", "TEC-10"], self.cfg)
        self.assertTrue(qs[0]["exec"])
        self.assertFalse(qs[1]["exec"])
        session = {"date": "2026-10-02", "day_number": 3, "level": 4, "questions": qs}
        # When the deck is built
        pptx, _ = build_deck(session, {"current": 2, "best": 2, "total": 2}, self.cfg, self.tmp)
        with zipfile.ZipFile(pptx) as z:
            exec_xml = z.read("ppt/slides/slide2.xml").decode()
        # Then the slide's bar runs for 25 seconds and is labelled "Headline window"
        self.assertIn('dur="25000"', exec_xml)
        self.assertIn('advTm="25000"', exec_xml)
        self.assertIn("Headline window", exec_xml)
        # And the hint reads "Headline → So what → What I need → Offer detail"
        self.assertIn(SHAPE, exec_xml)
        # And a constraint that needs more time than the window is dropped
        timed = Q._finish(dict(bank_question("EXE-07"), constraint="Finish in 45 seconds, then stop and hold the silence."),
                          "bank")
        asks = Q._finish(dict(bank_question("EXE-09"), constraint="End with a clear ask or next step."), "bank")
        Q.apply_exec_mode([timed, asks], self.cfg, dt.date(2026, 10, 2))
        self.assertEqual(timed["constraint"], "")
        self.assertEqual(asks["constraint"], "End with a clear ask or next step.")
        # (the plain question keeps its 60 seconds and no exec hint)
        text, _ = slides(str(pptx))[-2]
        self.assertIn(qs[1]["text"], text)
        with zipfile.ZipFile(pptx) as z:
            names = sorted(n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n))
            plain_xml = z.read(names[-2]).decode()
        self.assertIn('advTm="60000"', plain_xml)
        self.assertNotIn(SHAPE, plain_xml)

    def test_the_exec_says_thats_all_they_needed(self):
        # Given an exec slide whose reply is "Thanks, that's all I needed"
        q = with_reply(finished(["EXE-05"], self.cfg)[0], "thanks")
        session = {"date": "2026-10-02", "day_number": 3, "level": 3, "questions": [q]}
        pptx, _ = build_deck(session, {"current": 0, "best": 0, "total": 0}, self.cfg, self.tmp)
        shown = slides(str(pptx))
        # When the headline window ends
        reply_text, reply_notes = shown[2]  # intro, question, reply, closing
        # Then a reply slide shows that line for 10 seconds with no new question
        self.assertIn("Thanks, that’s all I needed.", reply_text)
        self.assertNotIn("OF 1", reply_text)
        self.assertNotIn(q["text"], reply_text)
        with zipfile.ZipFile(pptx) as z:
            self.assertIn('advTm="10000"', z.read("ppt/slides/slide3.xml").decode())
        # And the speaker notes ask "Did you stop?"
        self.assertIn("Did you stop?", reply_notes)

    def test_session_length_and_chapters_follow_the_exec_timeline(self):
        # Given a set with two exec questions, their replies and three 60-second questions
        qs = finished(["EXE-08", "TEC-10", "COM-10", "LEA-06", "DIS-03"], self.cfg)
        with_reply(qs[0], "thanks")
        with_reply(qs[2], "risk")
        self.assertEqual([q["exec"] for q in qs], [True, False, True, False, False])
        session = {"date": "2026-10-02", "day_number": 3, "level": 3, "questions": qs}
        # When the slide show is timed
        total = session_seconds(self.cfg, session)
        marks = chapters(session, self.cfg, offset=0)
        # Then its length counts each headline window and reply slide
        self.assertEqual(total, 10 + (25 + 10) + 60 + (25 + 20) + 60 + 60)
        self.assertIn("2 exec slides: 25-second headline", format_line(session, self.cfg))
        # And each YouTube chapter starts where its question slide starts
        self.assertEqual([t for t, _ in marks], [0, 10, 45, 105, 150, 210, 270])
        self.assertIn("to the CEO", marks[1][1])


# --------------------------------------------------------------------------- Feature: exec answer key

class AnswerKeyFeature(unittest.TestCase):
    feature = "answer_key.feature"

    def setUp(self):
        self.h = Home()

    def tearDown(self):
        self.h.cleanup()

    def session(self, date: str) -> dict:
        return self.h.history()["sessions"][date]

    def test_the_answer_key_arrives_after_recording(self):
        # Given today's five questions are recorded / When the practice session finishes
        self.h.run("daily", now="2026-10-01T05:30:00", answers="Start now,About right")
        ak = self.session("2026-10-01")["answer_key"]
        # Then an answer key page is saved with one card per question
        self.assertEqual(ak["status"], "ready")
        page = Path(ak["path"])
        self.assertEqual(len(cards(page)), 5)
        # And every card has a one-breath headline, the exec answer, the detail to offer and the trap
        for card in cards(page):
            for part in ("One breath", "The exec answer", "Hold back", "The trap", "Say it back"):
                self.assertIn(part, card)
            self.assertEqual(card.count('class="chip m-'), 4)
        # And the page opens once the practice chain is over
        self.assertTrue(ak["opened"])
        log = (self.h.dir / "home" / "logs" / "fde_coach.log").read_text(encoding="utf-8")
        opened = log.index(f"open  {page}")
        self.assertGreater(opened, log.index("[dry-run] YouTube upload"))

    def test_an_exec_reply_gets_its_own_model_answer(self):
        # Given an exec question whose listener replied "What's the biggest risk?"
        for kind in ("thanks", "need"):
            self.h.run("config", "--set", f"exec_mode.reply_weights.{kind}=0", now="2026-10-01T05:00:00")
        # When the answer key is written
        self.h.run("session", now="2026-10-01T06:00:00", answers="About right")
        s = self.session("2026-10-01")
        risky = [i for i, q in enumerate(s["questions"]) if (q.get("exec_reply") or {}).get("kind") == "risk"]
        self.assertTrue(risky)
        # Then its card shows the reply and a model answer to it
        card = cards(Path(s["answer_key"]["path"]))[risky[0]]
        self.assertIn("What’s the biggest risk?", card)
        self.assertIn("Without it by Friday, go-live moves a week.", card)

    def test_model_answers_stay_hidden_until_today_is_recorded(self):
        # Given today's deck is built but not recorded
        self.h.run("generate", now="2026-10-01T05:30:00")
        # When I ask for the answer key
        refused = self.h.run("answers", now="2026-10-01T05:40:00", check=False)
        # Then it is refused unless I pass --reveal
        self.assertNotEqual(refused.returncode, 0)
        self.assertIn("stay hidden until you record", refused.stderr)
        self.assertFalse((self.h.dir / "home" / "answers").exists() and
                         any((self.h.dir / "home" / "answers").glob("*.html")))
        shown = json.loads(self.h.run("answers", "--reveal", now="2026-10-01T05:41:00").stdout)
        self.assertEqual(shown["status"], "ready")

    def test_every_exec_answer_fits_the_headline_window(self):
        # Given Claude returns an exec answer longer than 70 words
        long_answer = {"slot": 1, "headline": "We are on track for the 15th with one risk.",
                       "answer": [{"move": "headline", "text": "We are on track. " * 6},
                                  {"move": "so_what", "text": "It matters because " + "the launch " * 20},
                                  {"move": "offer", "text": "Happy to go deeper if useful."}],
                       "detail": ["One.", "Two."], "trap": "Too much detail."}
        session = {"questions": [Q._finish(bank_question("EXE-04"), "bank")]}
        # When the answer key is validated
        ok, errors = A.validate([long_answer], session, DEFAULTS, [1])
        # Then that answer is rejected with the reason
        self.assertEqual(ok, {})
        self.assertIn("must be 20-70 words to fit the headline window", errors[0])
        # And the card shows the exec shape to fill in instead
        self.h.run("session", now="2026-10-01T06:00:00", answers="About right",
                   extra={"FAKE_CLAUDE_ANSWERS": "long"})
        ak = self.session("2026-10-01")["answer_key"]
        self.assertEqual(ak["status"], "pending")
        self.assertIn("20-70 words", ak["last_error"])
        page = cards(Path(ak["path"]))
        self.assertEqual(len(page), 5)
        self.assertTrue(all("Model answer pending" in c and 'class="moves fill"' in c for c in page))

    def test_offline_practice_still_gets_an_answer_key_later(self):
        # Given Claude is unreachable when the session finishes
        self.h.run("session", now="2026-10-01T06:00:00", answers="About right",
                   extra={"FAKE_CLAUDE_ANSWERS": "error"})
        ak = self.session("2026-10-01")["answer_key"]
        self.assertEqual(ak["status"], "pending")
        self.assertIn("Model answer pending", Path(ak["path"]).read_text(encoding="utf-8"))
        # When the 15-minute tick runs after Claude is back
        self.h.run("remind", now="2026-10-01T06:15:00")
        # Then the answer key is filled in
        ak = self.session("2026-10-01")["answer_key"]
        self.assertEqual(ak["status"], "ready")
        self.assertNotIn("Model answer pending", Path(ak["path"]).read_text(encoding="utf-8"))

    def test_every_one_breath_line_collects_on_one_page(self):
        # Given two recorded days with answer keys
        self.h.run("session", now="2026-10-01T06:00:00", answers="About right")
        self.h.run("session", now="2026-10-02T06:00:00", answers="About right")
        # When the index is built
        index = Path(self.h.run("answers", "--index", now="2026-10-02T07:00:00").stdout.strip())
        # Then it lists each day's one-breath lines, newest day first
        html = index.read_text(encoding="utf-8")
        self.assertLess(html.index("2026-10-02_Exec-Answer-Key.html"), html.index("2026-10-01_Exec-Answer-Key.html"))
        self.assertEqual(html.count("<li>"), 10)
        self.assertIn("I need one decision from you by Friday", html)


# --------------------------------------------------------------------------- the spec and the tests stay together

class TestFeatureFilesStayInSync(unittest.TestCase):
    def test_every_scenario_has_one_test_and_back(self):
        for cls in (ExecSlotsFeature, AnswerKeyFeature):
            wanted = {scenario_test_name(s) for s in scenarios(cls.feature)}
            have = {n for n in dir(cls) if n.startswith("test_")}
            self.assertEqual(wanted - have, set(), f"{cls.feature}: scenarios without a test")
            self.assertEqual(have - wanted, set(), f"{cls.__name__}: tests without a scenario")


class TestBankListeners(unittest.TestCase):
    def test_every_bank_question_names_its_listener(self):
        auds = Q.audiences()
        for q in Q.bank()["questions"]:
            self.assertIn(q["audience"], auds, q["id"])
            self.assertTrue(q["listener"], q["id"])
        counts = {a: sum(q["audience"] == a for q in Q.bank()["questions"]) for a in SENIOR}
        self.assertGreaterEqual(counts["exec"], 30)
        self.assertGreaterEqual(counts["account_leader"], 12)
        self.assertEqual(set(Q.bank()["exec_replies"]), {"thanks", "risk", "need"})

    def test_old_sessions_without_listeners_still_work(self):
        q = Q._finish({"category": "commercial", "format": "roleplay", "difficulty": 3,
                       "text": "The customer’s CFO says: “Give me one number.” Respond."}, "history")
        self.assertEqual((q["audience"], q["listener"]), ("exec", "Executive"))


if __name__ == "__main__":
    unittest.main()
